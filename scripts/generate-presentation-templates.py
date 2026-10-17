#!/usr/bin/env python3
"""Generate Open Office's original, editable presentation templates.

Requires python-pptx and fonttools. Each cover SVG and its PPTX use the same
drawing commands. SVG text uses paths from the app's fonts so previews do not
depend on system fonts. PPTX text stays editable. No downloaded artwork.
Run from any directory; outputs are deterministic ZIP packages.
"""

from __future__ import annotations

import hashlib
import html
import io
import json
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from lxml import etree
from fontTools.pens.svgPathPen import SVGPathPen
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from presentation_template_catalog import industry_templates


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "public" / "presentation-templates"
SANS = "WenQuanYi Micro Hei"
BOOK = "LXGW WenKai"
FONT_FILES = {
    SANS: ROOT / "assets/fonts/ui/WenQuanYiMicroHei-Regular.ttf",
    BOOK: ROOT / "assets/fonts/deck/LXGWWenKai-Regular.ttf",
}
FONTS = {family: TTFont(path) for family, path in FONT_FILES.items()}
WIDTH, HEIGHT = 1280, 720


def unit(px: float):
    return Inches(px / 96)


def color(value: str):
    return RGBColor.from_string(value.lstrip("#"))


class Canvas:
    def __init__(self, presentation: Presentation, background: str):
        self.slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self.svg: list[str] = []
        self.glyph_definitions: dict[tuple[str, str], tuple[str, str]] = {}
        self.rect(0, 0, WIDTH, HEIGHT, background)

    def shape(self, kind, x, y, w, h, fill, stroke=None, stroke_width=1):
        item = self.slide.shapes.add_shape(kind, unit(x), unit(y), unit(w), unit(h))
        if fill:
            item.fill.solid()
            item.fill.fore_color.rgb = color(fill)
        else:
            item.fill.background()
        if stroke:
            item.line.color.rgb = color(stroke)
            item.line.width = Pt(stroke_width * .75)
        else:
            item.line.fill.background()
        return item

    def rect(self, x, y, w, h, fill, stroke=None, stroke_width=1):
        self.shape(MSO_SHAPE.RECTANGLE, x, y, w, h, fill, stroke, stroke_width)
        self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{fill or "none"}" stroke="{stroke or "none"}" stroke-width="{stroke_width}"/>')

    def ellipse(self, x, y, w, h, fill, stroke=None, stroke_width=1):
        self.shape(MSO_SHAPE.OVAL, x, y, w, h, fill, stroke, stroke_width)
        self.svg.append(f'<ellipse cx="{x+w/2}" cy="{y+h/2}" rx="{w/2}" ry="{h/2}" fill="{fill or "none"}" stroke="{stroke or "none"}" stroke-width="{stroke_width}"/>')

    def line(self, x1, y1, x2, y2, stroke, width=1):
        item = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, unit(x1), unit(y1), unit(x2), unit(y2))
        item.line.color.rgb = color(stroke)
        item.line.width = Pt(width * .75)
        self.svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="{stroke}" stroke-width="{width}"/>')

    def text(self, x, y, w, text, size=24, fill="#172A46", bold=False,
             font=SANS, align="left", height=None, line_height=1.35):
        lines = text.split("\n")
        h = height or max(size * line_height * len(lines) + 8, size * 1.7)
        item = self.slide.shapes.add_textbox(unit(x), unit(y), unit(w), unit(h))
        frame = item.text_frame
        frame.clear()
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        for i, content in enumerate(lines):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = Pt(size * .75 * line_height)
            p.space_before = p.space_after = Pt(0)
            run = p.add_run()
            run.text = content
            run.font.name = font
            run.font.size = Pt(size * .75)
            run.font.bold = bold
            run.font.color.rgb = color(fill)
            # A run's East Asian typeface must be explicit, not inherited from
            # the package's default theme. The font itself is not embedded.
            east_asian = OxmlElement("a:ea")
            east_asian.set("typeface", font)
            run._r.get_or_add_rPr().append(east_asian)
            self.outlined_preview_text(x, y+i*size*line_height, w, content, size, fill, bold, font, align, line_height)
        return item

    def outlined_preview_text(self, x, y, w, content, size, fill, bold, family, align, line_height):
        """Use the same face, exact line pitch and synthetic bold as the engine."""
        font = FONTS[family]
        glyphs, cmap = font.getGlyphSet(), font.getBestCmap()
        upem = font["head"].unitsPerEm
        scale = size / upem
        positions, cursor = [], 0
        for char in content:
            name = cmap.get(ord(char), ".notdef")
            positions.append((name, cursor))
            cursor += font["hmtx"][name][0]
        text_width = cursor * scale
        origin = x + {"left": 0, "center": (w-text_width)/2, "right": w-text_width}[align]
        # The engine places exact line-spacing baselines at 80% of the pitch.
        baseline = y + size * line_height * .8
        shapes = []
        for name, advance in positions:
            key = (family, name)
            if key not in self.glyph_definitions:
                pen = SVGPathPen(glyphs)
                glyphs[name].draw(pen)
                ident = "glyph-" + hashlib.sha256((family+"/"+name).encode()).hexdigest()[:14]
                self.glyph_definitions[key] = (ident, pen.getCommands())
            ident, path = self.glyph_definitions[key]
            if not path:
                continue
            stroke = f' stroke="{fill}" stroke-width="{upem*.025}" stroke-linejoin="round"' if bold else ""
            shapes.append(f'<use href="#{ident}" transform="translate({origin+advance*scale:.5f} {baseline:.5f}) scale({scale:.8f} {-scale:.8f})"{stroke}/>')
        self.svg.append(f'<g fill="{fill}" aria-label="{html.escape(content, quote=True)}">'+"".join(shapes)+"</g>")

    def footer(self, name, page, fill="#617083", line="#DCE2E9"):
        self.line(72, 651, 1208, 651, line)
        self.text(72, 665, 800, f"OPEN OFFICE  /  {name}", 12, fill)
        self.text(1138, 665, 70, f"{page:02d} / 03", 12, fill, align="right")

    def preview(self, title):
        definitions = "".join(f'<path id="{ident}" d="{path}"/>' for ident, path in self.glyph_definitions.values() if path)
        return ('<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720" role="img">'
                f'<title>{html.escape(title)} · 封面预览</title><defs>{definitions}</defs>' + "".join(self.svg) + "</svg>\n")


def label(c, text, fill="#446384", x=72, y=58):
    c.text(x, y, 1000, text, 15, fill, True)


def headings(c, title="在这里填写内容标题", subtitle="用三个清晰的观点，讲述你的核心内容", fill="#172A46", x=72):
    c.text(x, 84, 1110, title, 46, fill, True)
    c.text(x, 153, 1110, subtitle, 21, fill)


def point(c, x, y, number, title, body, accent, ink, width=300):
    c.text(x, y, width, number, 18, accent, True)
    c.text(x, y + 40, width, title, 28, ink, True)
    c.text(x, y + 95, width, body, 19, ink, line_height=1.65)


def blueprint(p):
    navy, blue, pale = "#132C58", "#315FBE", "#E9EFFA"
    c = Canvas(p, "#F8FAFF")
    c.rect(848, 0, 432, 720, navy)
    for x, y, w, h, shade in [(894, 120, 170, 210, "#244B86"), (1016, 270, 220, 298, blue), (913, 432, 123, 137, "#78A4E8")]:
        c.rect(x, y, w, h, shade)
    c.line(890, 610, 1235, 610, "#8FB0E5", 2)
    label(c, "BUSINESS / STRATEGY", blue)
    c.rect(72, 180, 64, 7, blue)
    c.text(72, 218, 730, "让每个想法\n成为下一步", 68, navy, True)
    c.text(76, 427, 710, "在这里填写你的演示副标题", 25, "#5C6D86")
    c.text(76, 561, 710, "汇报人：你的名字  /  日期：20XX.XX", 18, navy)
    c.text(76, 658, 700, "OPEN OFFICE  ·  蓝图商务", 13, "#61738F")
    first = c
    c = Canvas(p, "#F8FAFF"); headings(c, fill=navy)
    for i, x in enumerate([72, 459, 846]):
        c.rect(x, 254, 362, 321, "#FFFFFF")
        c.rect(x, 254, 362, 6, blue if i != 1 else "#78A4E8")
        point(c, x + 28, 292, f"0{i+1}", ["明确方向", "制定计划", "推动结果"][i], "在这里输入关键观点\n补充事实、数据或案例\n让下一步行动更加清楚", blue, navy)
    c.footer("蓝图商务", 2)
    c = Canvas(p, navy); c.rect(0, 0, 25, 720, "#78A4E8")
    c.text(80, 166, 1120, "把想法，变成行动", 56, "#FFFFFF", True)
    c.text(80, 275, 1120, "Open Office", 94, "#FFFFFF", True)
    c.line(84, 447, 1210, 447, "#5C7EB3", 2)
    c.text(84, 484, 1100, "谢谢观看  /  期待与你交流", 25, "#BFD4F4")
    return first


def midnight_tech(p):
    bg, teal, white = "#101C2D", "#56E0C0", "#EDF7FF"
    c = Canvas(p, bg)
    for radius, shade in [(215, "#233A51"), (160, "#315B6B"), (100, teal)]:
        c.ellipse(1000-radius, 351-radius, radius*2, radius*2, None, shade, 2)
    for x, y in [(1000, 136), (841, 351), (1100, 351), (1000, 511)]:
        c.ellipse(x-8, y-8, 16, 16, teal)
    c.line(785, 351, 1215, 351, "#294455")
    c.line(1000, 100, 1000, 598, "#294455")
    label(c, "IDEAS / SYSTEMS / FUTURE", teal)
    c.text(72, 198, 700, "探索新的可能", 68, white, True)
    c.text(76, 311, 680, "用清晰的思考，连接未来", 31, teal)
    c.text(76, 414, 680, "在这里填写主题简介或技术副标题", 23, "#A5BBCD")
    c.text(76, 567, 680, "分享人：你的名字  /  20XX", 18, "#A5BBCD")
    c.footer("深空科技", 1, "#7996AF", "#294455")
    first = c
    c = Canvas(p, bg); headings(c, fill=white, subtitle="从问题出发，用结构化思考找到答案")
    for i, y in enumerate([246, 371, 496]):
        c.rect(72, y, 1136, 104, "#172A3F")
        c.rect(72, y, 5, 104, teal)
        c.text(105, y+22, 75, f"0{i+1}", 35, teal, True)
        c.text(222, y+20, 280, ["问题与洞察", "方法与实现", "成果与展望"][i], 27, white, True)
        c.text(570, y+25, 590, "在这里输入核心内容，让逻辑清晰可见。", 21, "#A5BBCD")
    c.footer("深空科技", 2, "#7996AF", "#294455")
    c = Canvas(p, bg)
    for radius, shade in [(264, "#233A51"), (217, "#315B6B"), (170, "#397D83")]:
        c.ellipse(640-radius, 360-radius, radius*2, radius*2, None, shade, 2)
    c.rect(200, 274, 880, 170, bg)
    c.text(200, 278, 880, "Open Office", 88, white, True, align="center")
    c.text(200, 407, 880, "让下一次探索，从这里开始", 25, teal, align="center")
    return first


def ivory_study(p):
    bg, ink, tan = "#F4F0E5", "#3A3A32", "#A28C60"
    c = Canvas(p, bg)
    label(c, "RESEARCH / PERSPECTIVE", tan)
    c.line(72, 118, 1208, 118, tan)
    c.text(72, 182, 760, "在思考中\n找到新的答案", 65, ink, True, BOOK)
    c.text(77, 394, 740, "研究主题、课程分享与观点表达", 25, "#746F61")
    c.text(77, 551, 720, "作者：你的名字  /  单位或院系", 18, ink)
    c.rect(943, 213, 59, 325, "#D8CDB6")
    c.rect(1021, 300, 59, 238, "#BCAF8E")
    c.rect(1099, 383, 59, 155, tan)
    c.ellipse(945, 159, 210, 210, None, tan, 2)
    c.footer("暖白学术", 1, "#837B69", "#D7CEBA")
    first = c
    c = Canvas(p, bg)
    c.text(72, 70, 1120, "在这里填写内容标题", 48, ink, True, BOOK)
    c.line(72, 159, 1208, 159, tan)
    c.text(72, 218, 440, "从一个好问题开始", 32, ink, True, BOOK)
    c.text(72, 284, 430, "在这里说明研究背景、讨论范围\n或你希望听众思考的关键问题。\n\n为观点提供依据，让表达更有力。", 22, "#746F61", line_height=1.75)
    for i, y in enumerate([220, 348, 476]):
        c.text(610, y, 70, f"0{i+1}", 31, tan, True, BOOK)
        c.text(708, y, 485, ["背景与问题", "分析与论证", "结论与启发"][i], 27, ink, True)
        c.text(708, y+48, 485, "在这里填写支撑观点的简短说明。", 20, "#746F61")
        if i < 2: c.line(610, y+96, 1208, y+96, "#D7CEBA")
    c.footer("暖白学术", 2, "#837B69", "#D7CEBA")
    c = Canvas(p, bg)
    c.line(170, 164, 1110, 164, tan)
    c.text(170, 251, 940, "谢谢观看", 65, ink, True, BOOK, "center")
    c.text(170, 366, 940, "Open Office", 44, tan, False, BOOK, "center")
    c.text(170, 463, 940, "保持好奇，继续思考", 23, "#746F61", align="center")
    c.line(170, 553, 1110, 553, tan)
    return first


def forest_notes(p):
    green, bg, mint = "#214D3E", "#F2F5EC", "#A7C7A3"
    c = Canvas(p, bg)
    c.rect(824, 0, 456, 720, green)
    for x, y, w, h, fill in [(882, 100, 150, 257, mint), (1038, 270, 175, 289, "#7DAB83"), (918, 432, 100, 170, "#C5D8B6")]:
        c.ellipse(x, y, w, h, fill)
    c.line(1054, 189, 1054, 626, "#DDE7D3", 2)
    c.line(953, 275, 1054, 389, "#DDE7D3", 2)
    c.line(1133, 420, 1054, 510, "#DDE7D3", 2)
    label(c, "NATURE / GROWTH / BALANCE", green)
    c.text(72, 212, 700, "让灵感\n自然生长", 72, green, True)
    c.text(77, 432, 680, "在这里填写你的演示副标题", 25, "#6C806C")
    c.text(77, 560, 690, "分享人：你的名字  /  20XX", 18, green)
    c.text(77, 661, 650, "OPEN OFFICE  ·  森林随想", 13, "#6C806C")
    first = c
    c = Canvas(p, bg); headings(c, fill=green, subtitle="关注长期价值，让每一个行动有所生长")
    for i, x in enumerate([72, 467, 862]):
        c.ellipse(x, 248, 62, 88, [green, "#7DAB83", mint][i])
        c.line(x+32, 338, x+32, 593, "#BCD0B4", 2)
        point(c, x+89, 270, f"0{i+1}", ["观察与发现", "行动与改变", "持续与成长"][i], "写下你的核心观点\n补充一个具体案例\n给出可执行的建议", green, green, 250)
    c.footer("森林随想", 2, "#6C806C", "#D3DECA")
    c = Canvas(p, green)
    c.ellipse(106, 96, 184, 365, "#2B5D49")
    c.ellipse(970, 252, 208, 375, "#34684F")
    c.text(250, 252, 780, "Open Office", 88, "#EDF4E8", True, align="center")
    c.text(250, 391, 780, "愿每个想法，都能生根发芽", 26, "#BBD2AE", align="center")
    return first


def coral_studio(p):
    coral, ink, cream = "#E96B57", "#692F3B", "#FFF4E8"
    c = Canvas(p, coral)
    c.ellipse(938, 112, 259, 259, cream)
    c.rect(843, 316, 234, 234, ink)
    c.ellipse(1030, 465, 140, 140, "#F6BA85")
    label(c, "CREATIVE STUDIO / IDEAS IN MOTION", cream)
    c.text(68, 187, 790, "让创意\n被看见", 92, cream, True)
    c.text(77, 449, 715, "关于灵感、表达与新的可能", 27, cream)
    c.text(77, 576, 715, "创作者：你的名字", 19, cream)
    c.text(77, 662, 715, "OPEN OFFICE  ·  珊瑚创意", 13, cream)
    first = c
    c = Canvas(p, cream); headings(c, fill=ink, subtitle="三个想法，开启一次新的创作")
    for i, x in enumerate([72, 470, 868]):
        c.text(x, 229, 300, f"0{i+1}", 100, coral, True)
        c.line(x, 377, x+336, 377, "#E4C9B7", 2)
        c.text(x, 407, 336, ["灵感从哪里来", "怎样表达想法", "让作品发生"][i], 28, ink, True)
        c.text(x, 471, 336, "在这里写下你的观点\n用简单的语言讲述故事\n留一点空间，给新的灵感", 20, ink, line_height=1.6)
    c.footer("珊瑚创意", 2, "#9B6867", "#E4C9B7")
    c = Canvas(p, ink)
    c.ellipse(302, 48, 676, 676, coral)
    c.text(160, 239, 960, "Open Office", 95, cream, True, align="center")
    c.text(160, 389, 960, "你的下一次创作，从这里开始", 26, cream, align="center")
    return first


def gradient(c, start, end):
    a = tuple(int(start[i:i+2], 16) for i in (1, 3, 5))
    b = tuple(int(end[i:i+2], 16) for i in (1, 3, 5))
    # Editable adjacent shapes form the exact same soft gradient in both files.
    for i in range(64):
        shade = "#" + "".join(f"{round(a[j]+(b[j]-a[j])*i/63):02X}" for j in range(3))
        c.rect(i*20, 0, 20, 720, shade)


def violet_dawn(p):
    ink, purple = "#47356B", "#8461BA"
    c = Canvas(p, "#F0E9FF"); gradient(c, "#F0E9FF", "#FBE5E0")
    c.ellipse(905, 138, 250, 250, "#C9B4E9")
    c.ellipse(1004, 333, 170, 170, "#EDC0CD")
    c.ellipse(831, 349, 190, 190, None, purple, 3)
    label(c, "VISION / POSSIBILITY", purple)
    c.text(72, 218, 750, "每个想法\n都有新的光芒", 66, ink, True)
    c.text(77, 424, 720, "在这里填写你的主题与副标题", 25, "#79658E")
    c.text(77, 564, 720, "分享人：你的名字  /  20XX", 18, ink)
    c.text(77, 662, 720, "OPEN OFFICE  ·  紫霞愿景", 13, "#79658E")
    first = c
    c = Canvas(p, "#F6F1FC"); headings(c, fill=ink, subtitle="把宏大的愿景，拆解成清晰的表达")
    for i, x in enumerate([72, 459, 846]):
        c.rect(x, 263, 362, 316, ["#EAE0F8", "#EEE3F5", "#F5E3EC"][i])
        c.ellipse(x+25, 291, 48, 48, [purple, "#A382C1", "#BC8CA6"][i])
        c.text(x+25, 296, 48, str(i+1), 24, "#FFFFFF", True, align="center")
        c.text(x+28, 372, 310, ["描绘愿景", "聚焦价值", "一起实现"][i], 30, ink, True)
        c.text(x+28, 437, 310, "在这里填写内容要点\n为听众呈现清晰的方向\n让想法连接更多可能", 20, "#79658E", line_height=1.6)
    c.footer("紫霞愿景", 2, "#79658E", "#DECEE9")
    c = Canvas(p, "#F0E9FF"); gradient(c, "#E3D5F6", "#F7D9E0")
    c.ellipse(80, 111, 185, 185, "#D4BCEB"); c.ellipse(1040, 447, 140, 140, "#EDC0CD")
    c.text(180, 246, 920, "Open Office", 92, ink, True, align="center")
    c.text(180, 393, 920, "让想法发光，让未来发生", 27, "#79658E", align="center")
    return first


def mono_editorial(p):
    black, bg = "#20201E", "#F5F4EF"
    c = Canvas(p, bg)
    label(c, "OPEN OFFICE / EDITORIAL", black)
    c.text(930, 38, 278, "ISSUE 01", 15, black, True, align="right")
    c.line(72, 102, 1208, 102, black, 2)
    c.rect(900, 145, 308, 428, black)
    c.text(919, 192, 270, "01", 145, bg, True, align="center")
    c.text(922, 439, 264, "A NEW\nPERSPECTIVE", 23, bg, True, align="center")
    c.text(72, 174, 800, "观点，\n值得被听见。", 74, black, True, BOOK)
    c.text(77, 418, 750, "在这里填写主题、观点或故事副标题", 24, "#666660")
    c.text(77, 551, 750, "撰稿 / 分享：你的名字", 18, black)
    c.footer("黑白编辑部", 1, black, black)
    first = c
    c = Canvas(p, bg)
    c.text(72, 65, 1136, "在这里填写内容标题", 47, black, True, BOOK)
    c.line(72, 157, 1208, 157, black, 3)
    c.line(619, 204, 619, 605, "#B7B7AE")
    c.text(72, 213, 497, "01", 91, black, True)
    c.text(72, 341, 497, "把核心观点放在这里", 32, black, True, BOOK)
    c.text(72, 410, 497, "在这里讲述最重要的观点。\n用事实建立联系，用细节让内容生动。\n\n保持表达的节奏，给思考留出空间。", 22, "#666660", line_height=1.65)
    for i, y in enumerate([221, 427]):
        c.text(675, y, 60, f"0{i+2}", 25, black, True)
        c.text(774, y, 434, ["另一种观察", "进一步的思考"][i], 29, black, True, BOOK)
        c.text(774, y+63, 434, "在这里补充一个角度、一个案例\n或值得继续讨论的问题。", 21, "#666660", line_height=1.6)
    c.line(675, 385, 1208, 385, "#B7B7AE")
    c.footer("黑白编辑部", 2, black, black)
    c = Canvas(p, bg)
    c.rect(72, 72, 1136, 30, black)
    c.text(72, 201, 1136, "Open Office.", 120, black, True)
    c.line(72, 411, 1208, 411, black, 2)
    c.text(77, 460, 1126, "谢谢观看。下一个观点，等你来写。", 29, black, False, BOOK)
    return first


def retro_warmth(p):
    bg, brown, orange = "#F3E5CB", "#65402E", "#BC6942"
    c = Canvas(p, bg)
    c.rect(38, 38, 1204, 644, None, brown, 2)
    c.rect(50, 50, 1180, 620, None, "#C7AB82")
    c.ellipse(901, 174, 230, 230, orange)
    c.rect(901, 289, 230, 258, orange)
    c.ellipse(942, 237, 148, 148, bg)
    c.rect(942, 311, 148, 236, bg)
    c.line(875, 548, 1158, 548, brown, 3)
    label(c, "STORIES / MOMENTS / CONNECTIONS", brown, 88, 83)
    c.text(88, 206, 740, "把故事\n慢慢讲给你听", 64, brown, True, BOOK)
    c.text(93, 414, 725, "在这里写下你的故事副标题", 25, "#987453")
    c.text(93, 555, 725, "讲述人：你的名字", 19, brown)
    c.text(93, 618, 725, "OPEN OFFICE  ·  暖调复古", 13, "#987453")
    first = c
    c = Canvas(p, bg)
    c.text(72, 72, 1136, "在这里填写内容标题", 48, brown, True, BOOK)
    c.text(77, 151, 1110, "用一条时间线，串联你想讲述的故事", 23, "#987453")
    c.line(117, 300, 1133, 300, "#C7AB82", 3)
    for i, x in enumerate([72, 459, 846]):
        c.ellipse(x, 255, 90, 90, orange if i == 1 else brown)
        c.text(x, 272, 90, f"0{i+1}", 30, bg, True, BOOK, "center")
        c.text(x, 389, 335, ["故事的起点", "路上的发现", "新的篇章"][i], 31, brown, True, BOOK)
        c.text(x, 458, 335, "写下一段关键的经历\n留住一个有温度的细节\n让每一步都值得被记住", 21, "#987453", line_height=1.6)
    c.footer("暖调复古", 2, "#987453", "#C7AB82")
    c = Canvas(p, orange)
    c.rect(42, 42, 1196, 636, None, bg, 2)
    c.text(160, 229, 960, "Open Office", 96, bg, True, BOOK, "center")
    c.text(160, 390, 960, "故事未完，我们下次再见", 27, bg, align="center")
    return first


def electric_launch(p):
    yellow, black = "#E5EF65", "#20261D"
    c = Canvas(p, yellow)
    c.rect(1004, 0, 276, 720, black)
    c.ellipse(1074, 129, 140, 140, None, yellow, 12)
    c.rect(1074, 339, 140, 140, yellow)
    c.ellipse(1100, 534, 88, 88, yellow)
    label(c, "LAUNCH / MAKE IT HAPPEN", black)
    c.text(69, 166, 900, "下一站，\n现在出发！", 88, black, True)
    c.rect(77, 447, 550, 58, black)
    c.text(99, 457, 510, "你的发布主题与一句话主张", 24, yellow, True)
    c.text(77, 576, 850, "发布人：你的名字  /  20XX", 20, black)
    c.text(77, 662, 850, "OPEN OFFICE  ·  活力发布", 13, black)
    first = c
    c = Canvas(p, "#F6F8E8"); headings(c, fill=black, subtitle="把亮点说清楚，把行动放在眼前")
    for i, (x, y, h) in enumerate([(72, 351, 250), (459, 290, 311), (846, 229, 372)]):
        c.rect(x, y, 362, h, ["#E5EF65", "#CBDC76", black][i])
        ink = yellow if i == 2 else black
        c.text(x+27, y+22, 300, f"0{i+1}", 40, ink, True)
        c.text(x+27, y+92, 310, ["看见亮点", "理解价值", "马上行动"][i], 32, ink, True)
        c.text(x+27, y+158, 310, "在这里填写你的关键内容\n让信息简洁、有力、直接", 20, ink, line_height=1.6)
    c.footer("活力发布", 2, "#6E7959", "#CAD3AF")
    c = Canvas(p, black)
    c.text(72, 137, 1136, "LET’S GO.", 112, yellow, True)
    c.text(79, 326, 1110, "Open Office", 82, "#F6F8E8", True)
    c.line(80, 474, 1200, 474, "#778340", 2)
    c.text(80, 526, 1110, "让下一步，从现在开始。", 31, yellow, True)
    return first


def clear_space(p):
    bg, teal, ink = "#F2F8F9", "#438B92", "#29434E"
    c = Canvas(p, bg)
    c.line(72, 73, 1208, 73, "#CADDE0")
    c.ellipse(559, 134, 162, 162, None, "#C5DEE2", 2)
    c.ellipse(603, 178, 74, 74, None, teal, 2)
    c.text(160, 322, 960, "清晰，自有力量", 61, ink, False, align="center")
    c.text(160, 423, 960, "在这里填写你的主题与副标题", 24, "#77929B", align="center")
    c.text(160, 555, 960, "分享人：你的名字", 18, teal, align="center")
    c.line(72, 648, 1208, 648, "#CADDE0")
    c.text(72, 666, 1136, "OPEN OFFICE  ·  清透留白", 12, "#77929B", align="center")
    first = c
    c = Canvas(p, bg)
    c.rect(72, 83, 6, 52, teal)
    c.text(108, 76, 1085, "在这里填写内容标题", 45, ink)
    for i, y in enumerate([231, 365, 499]):
        c.text(110, y, 85, f"0{i+1}", 22, teal)
        c.text(246, y-5, 340, ["一句核心观点", "一个具体例子", "一项明确行动"][i], 29, ink)
        c.text(700, y+1, 500, "在这里补充简洁、清晰的内容说明。", 21, "#77929B")
        if i < 2: c.line(110, y+84, 1208, y+84, "#CADDE0")
    c.footer("清透留白", 2, "#77929B", "#CADDE0")
    c = Canvas(p, bg)
    c.line(640, 96, 640, 203, teal, 2)
    c.text(160, 276, 960, "Open Office", 84, ink, False, align="center")
    c.text(160, 414, 960, "谢谢观看", 25, teal, align="center")
    c.line(640, 513, 640, 620, teal, 2)
    return first


TEMPLATES = [
    ("blueprint", "蓝图商务", "蓝白分栏与建筑几何，适合项目汇报和商业方案。", "#315FBE", blueprint),
    ("midnight-tech", "深空科技", "深色轨道与薄荷亮色，适合技术分享和产品研究。", "#56E0C0", midnight_tech),
    ("ivory-study", "暖白学术", "纸感暖白与书卷气标题，适合课程、研究和观点分享。", "#A28C60", ivory_study),
    ("forest-notes", "森林随想", "森林绿与植物几何，适合自然、生活和长期规划。", "#214D3E", forest_notes),
    ("coral-studio", "珊瑚创意", "珊瑚橘与大胆拼贴，适合创意提案和作品展示。", "#E96B57", coral_studio),
    ("violet-dawn", "紫霞愿景", "柔和渐变与圆形构成，适合愿景、品牌和灵感分享。", "#8461BA", violet_dawn),
    ("mono-editorial", "黑白编辑部", "黑白对比与报刊分栏，适合故事、评论和设计表达。", "#20201E", mono_editorial),
    ("retro-warmth", "暖调复古", "陶土暖色与双线边框，适合文化、经历和故事讲述。", "#BC6942", retro_warmth),
    ("electric-launch", "活力发布", "荧光黄与阶梯版式，适合产品发布和行动倡议。", "#A9BE31", electric_launch),
    ("clear-space", "清透留白", "清透蓝绿与充足留白，适合简洁汇报和专注表达。", "#438B92", clear_space),
]
TEMPLATES += industry_templates(Canvas)


def canonical_package(raw: bytes) -> bytes:
    """Remove irrelevant seed metadata and use stable ZIP timestamps/order."""
    output = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as source, ZipFile(output, "w", ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            if name == "docProps/thumbnail.jpeg" or name.startswith("ppt/printerSettings/"):
                continue
            data = source.read(name)
            if name == "docProps/app.xml":
                data = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
                        '<Application>Open Office</Application><PresentationFormat>Widescreen</PresentationFormat>'
                        '<Slides>3</Slides><Notes>0</Notes><HiddenSlides>0</HiddenSlides><AppVersion>1.0</AppVersion>'
                        '</Properties>').encode()
            elif name.endswith(".rels") or name == "[Content_Types].xml":
                root = etree.fromstring(data)
                for item in list(root):
                    kind = item.get("Type", "")
                    if kind.endswith("/thumbnail") or kind.endswith("/printerSettings") or item.get("Extension") in {"bin", "jpeg"}:
                        root.remove(item)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif name.startswith("ppt/theme/") and name.endswith(".xml"):
                root = etree.fromstring(data)
                root.set("name", "Open Office")
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                for item in root.findall(".//a:fontScheme", ns):
                    item.set("name", "Open Office")
                    for family in item.iter():
                        if family.get("typeface") is not None:
                            family.set("typeface", SANS)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                root = etree.fromstring(data)
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
                # add_shape supplies a default theme effect reference. All
                # template colors/strokes are explicit: never inherit shadows.
                for style in root.findall(".//p:style", ns):
                    style.getparent().remove(style)
                for properties in root.findall(".//p:spPr", ns):
                    for effect in properties.findall("a:effectLst", ns):
                        properties.remove(effect)
                    etree.SubElement(properties, "{"+ns["a"]+"}effectLst")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            info = ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            target.writestr(info, data)
    return output.getvalue()


def verify(path: Path):
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with ZipFile(path) as package:
        slides = [n for n in package.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        assert len(slides) == 3, f"{path.name}: expected three slides"
        counts = []
        for name in slides:
            root = ET.fromstring(package.read(name))
            texts = root.findall(".//a:t", ns)
            assert len(texts) >= 2, f"{path.name}: missing editable text on {name}"
            counts.append(len(texts))
        for name in package.namelist():
            if name.endswith(".rels"):
                root = ET.fromstring(package.read(name))
                assert all(item.get("TargetMode") != "External" for item in root), f"{path.name}: external relationship"
        assert not any(n.startswith("ppt/media/") for n in package.namelist()), "Templates must contain editable vector shapes"
    return {"file": path.name, "slides": 3, "editableTextRuns": counts, "bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def main():
    assert len(TEMPLATES) == 100, "The built-in gallery must contain exactly 100 templates"
    assert len({entry[0] for entry in TEMPLATES}) == len(TEMPLATES), "Template IDs must be unique"
    assert len({entry[1] for entry in TEMPLATES}) == len(TEMPLATES), "Template names must be unique"
    OUT.mkdir(parents=True, exist_ok=True)
    manifest, reports = [], []
    for ident, name, description, accent, draw in TEMPLATES:
        p = Presentation()
        p.slide_width, p.slide_height = unit(WIDTH), unit(HEIGHT)
        p._element.sldSz.set("type", "screen16x9")
        p.core_properties.title = name
        p.core_properties.subject = "Open Office 原创演示模板"
        p.core_properties.author = "Open Office"
        p.core_properties.last_modified_by = "Open Office"
        p.core_properties.created = p.core_properties.modified = datetime(2026, 1, 1)
        p.core_properties.comments = "Original editable vector artwork; generated by generate-presentation-templates.py."
        cover = draw(p)
        raw = io.BytesIO(); p.save(raw)
        file, preview = f"{ident}.pptx", f"{ident}.svg"
        (OUT / file).write_bytes(canonical_package(raw.getvalue()))
        (OUT / preview).write_text(cover.preview(name), encoding="utf-8")
        manifest.append({"id": ident, "name": name, "description": description, "accent": accent, "slideCount": 3, "file": file, "preview": preview})
        reports.append(verify(OUT / file))
    (OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps({"templates": len(reports), "slides": sum(r["slides"] for r in reports), "totalBytes": sum(r["bytes"] for r in reports), "checks": reports}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
