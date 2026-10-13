#!/usr/bin/env python3
"""Independent lxml/python-pptx package and preservation audit. SPDX-License-Identifier: MIT.

Usage: inspect.py SOURCE [OUTPUT ...] --fixture ID --scenario noop|edit|duplicate|format-color|format-size|duplicate-object --report FILE
No application modules, browser or Office process are used. XML comparisons use
expanded names and resolved relationships, including actual motion paths, trigger
and connector targets, text runs/fields/styles, fills and image crop/alpha data.
"""
from __future__ import annotations
import os, sys
# Avoid shadowing Python's standard inspect module with this CLI's filename.
sys.path = [entry for entry in sys.path if os.path.abspath(entry) != os.path.dirname(os.path.abspath(__file__))]
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
import argparse,copy,fnmatch,hashlib,json,math,posixpath,sys,urllib.parse
from lxml import etree as E
from pptx import Presentation

ROOT=Path(__file__).resolve().parent
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
SHAPES={'sp','pic','graphicFrame','grpSp','cxnSp','AlternateContent'}
PARSER=E.XMLParser(resolve_entities=False,no_network=True,remove_blank_text=True)
def local(n):return E.QName(n).localname

def xml(data):
    if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():raise ValueError('DTD/entity declaration forbidden')
    return E.fromstring(data,PARSER)
def relpart(part):return posixpath.join(posixpath.dirname(part),'_rels',posixpath.basename(part)+'.rels') if part else '_rels/.rels'
def resolve(owner,target):
    path=urllib.parse.unquote(urllib.parse.urlsplit(target).path)
    result=posixpath.normpath(path.lstrip('/') if path.startswith('/') else posixpath.join(posixpath.dirname(owner),path))
    if result.startswith('../') or result=='..':raise ValueError('relationship escapes archive')
    return result
def first(node,path):
    values=node.xpath(path,namespaces=NS) if node is not None else []
    return values[0] if values else None
def shape_name(node):
    meta=first(node,'.//p:cNvPr')
    return meta.get('name','') if meta is not None else local(node)
def transform_node(shape):return first(shape,'./p:xfrm|./p:spPr/a:xfrm|./p:grpSpPr/a:xfrm')
def notes_text(pkg,part):
    target=pkg.target(part,'notesSlide')
    if not target:return ''
    doc=pkg.docs[target]
    bodies=doc.xpath('.//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]/p:txBody',namespaces=NS)
    return '\n'.join('\n'.join(''.join(p.xpath('.//a:t/text()',namespaces=NS)) for p in body.xpath('./a:p',namespaces=NS)) for body in bodies)

class Package:
    def __init__(self,path):
        self.path=str(path);self.parts={};self.docs={};self.rels={};self.errors=[];self.internal=0;self.external=[]
        raw=Path(path).read_bytes() if not isinstance(path,bytes) else path
        self.sha256=hashlib.sha256(raw).hexdigest()
        with ZipFile(BytesIO(raw)) as archive:
            bad=archive.testzip()
            if bad:self.errors.append({'kind':'zip_crc','part':bad})
            for name in archive.namelist():
                if name in self.parts:self.errors.append({'kind':'duplicate_zip_part','part':name})
                self.parts[name]=archive.read(name)
        for name,data in self.parts.items():
            if name.endswith(('.xml','.rels')):
                try:self.docs[name]=xml(data)
                except Exception as error:self.errors.append({'kind':'xml_invalid','part':name,'message':str(error)})
        for name,doc in self.docs.items():
            if not name.endswith('.rels'):continue
            folder,base=posixpath.split(name);owner='' if name=='_rels/.rels' else posixpath.join(posixpath.dirname(folder),base[:-5]);relations={}
            for relation in doc:
                rid=relation.get('Id');target=relation.get('Target','');external=relation.get('TargetMode')=='External'
                if rid in relations:self.errors.append({'kind':'duplicate_relationship','part':name,'id':rid})
                try:resolved=target if external else resolve(owner,target)
                except Exception as error:self.errors.append({'kind':'unsafe_target','part':name,'target':target,'message':str(error)});continue
                relations[rid]={'type':relation.get('Type',''),'target':resolved,'external':external}
                if external:self.external.append({'owner':owner,'target':target})
                else:
                    self.internal+=1
                    if resolved not in self.parts:self.errors.append({'kind':'missing_internal_target','part':name,'id':rid,'target':resolved})
            self.rels[owner]=relations
        for part,doc in self.docs.items():
            if part.endswith('.rels'):continue
            for node in doc.iter():
                for prefix in ([*node.get('Requires','').split()] if local(node)=='Choice' else []) + node.get('{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable','').split():
                    if prefix not in node.nsmap:self.errors.append({'kind':'undeclared_mc_requires_prefix','part':part,'prefix':prefix})
                for key,value in node.attrib.items():
                    if key in ['{'+NS['r']+'}'+name for name in ['id','embed','link']] and value and value not in self.rels.get(part,{}):self.errors.append({'kind':'unresolved_relationship_id','part':part,'id':value})
        pres=self.target('','officeDocument') or 'ppt/presentation.xml';self.presentation=pres
        doc=self.docs[pres];self.slides=[];self.slideids=[]
        for entry in doc.xpath('./p:sldIdLst/p:sldId',namespaces=NS):
            self.slides.append(self.rels[pres][entry.get('{'+NS['r']+'}id')]['target']);self.slideids.append(entry.get('id'))
        for part in self.slides:self.check_slide(part)
        self.readability=self.python_read(raw)
        self.embedded_workbooks=[]
        for name,data in self.parts.items():
            if name.endswith('.xlsx'):
                try:
                    with ZipFile(BytesIO(data)) as workbook:
                        members=set(workbook.namelist());missing=[]
                        for member in members:
                            if not member.endswith('.rels'):continue
                            folder,base=posixpath.split(member);owner='' if member=='_rels/.rels' else posixpath.join(posixpath.dirname(folder),base[:-5])
                            for relation in xml(workbook.read(member)):
                                if relation.get('TargetMode')!='External' and resolve(owner,relation.get('Target','')) not in members:missing.append(relation.get('Target'))
                        self.embedded_workbooks.append({'part':name,'entries':len(members),'missingTargets':missing})
                        if missing:self.errors.append({'kind':'xlsx_internal_target','part':name,'targets':missing})
                except Exception as error:self.errors.append({'kind':'xlsx_invalid','part':name,'message':str(error)})
    def target(self,owner,kind):return next((r['target'] for r in self.rels.get(owner,{}).values() if r['type'].endswith('/'+kind) and not r['external']),None)
    def check_slide(self,part):
        doc=self.docs[part];metas=doc.xpath('.//p:cNvPr',namespaces=NS);ids={n.get('id') for n in metas}
        def exclusive(a,b):
            def branches(node):
                return {ancestor.getparent():ancestor for ancestor in node.iterancestors() if local(ancestor) in ['Choice','Fallback'] and ancestor.getparent() is not None and local(ancestor.getparent())=='AlternateContent'}
            aa,bb=branches(a),branches(b)
            return any(parent in bb and branch is not bb[parent] for parent,branch in aa.items())
        for i,a in enumerate(metas):
            for b in metas[i+1:]:
                if a.get('id')==b.get('id') and not exclusive(a,b):self.errors.append({'kind':'duplicate_shape_id','part':part,'id':a.get('id')})
        timing_ids=[n.get('id') for n in doc.xpath('.//p:cTn',namespaces=NS)]
        if len(timing_ids)!=len(set(timing_ids)):self.errors.append({'kind':'duplicate_timing_id','part':part})
        for node in doc.iter():
            target=node.get('spid') or (node.get('id') if local(node) in ['stCxn','endCxn'] else None)
            if target and target not in ids:self.errors.append({'kind':'dangling_shape_reference','part':part,'tag':local(node),'target':target})
            if local(node)=='tn' and node.get('val') not in timing_ids:self.errors.append({'kind':'dangling_timing_reference','part':part,'target':node.get('val')})
    def python_read(self,raw):
        try:
            deck=Presentation(BytesIO(raw));slides=[]
            for slide in deck.slides:
                tables=[];charts=[]
                def walk(shapes):
                    for shape in shapes:
                        if shape.has_table:tables.append([[cell.text for cell in row.cells] for row in shape.table.rows])
                        if shape.has_chart:
                            chart=shape.chart;charts.append({'plotTypes':[type(p).__name__ for p in chart.plots],'series':[{'name':series.name,'values':list(series.values)} for series in chart.series]})
                        if hasattr(shape,'shapes'):walk(shape.shapes)
                walk(slide.shapes);slides.append({'tables':tables,'charts':charts})
            return {'readable':True,'pages':len(deck.slides),'slides':slides}
        except Exception as error:return {'readable':False,'error':str(error)}
    def shape_map(self,part):return {n.get('id'):n.get('name','') for n in self.docs[part].xpath('.//p:cNvPr',namespaces=NS)}
    def summary(self):return {'path':self.path,'sha256':self.sha256,'zipEntries':len(self.parts),'slides':len(self.slides),'internalRelationships':self.internal,'externalRelationships':self.external,'errors':self.errors,'pythonPptx':self.readability,'embeddedWorkbooks':self.embedded_workbooks}

def semantic(node,pkg,part,format_mode=None,format_target=None):
    if node is None:return None
    tag=local(node);parent=local(node.getparent()) if node.getparent() is not None else ''
    if format_target is not None and tag in SHAPES and shape_name(node)==format_target:format_target=None
    active_format=format_mode if format_target is None else None
    if tag=='xfrm':return None # geometry is independently resolved below, including group transforms and inherited off/ext.
    if active_format=='format-color' and tag=='solidFill' and parent in ['rPr','defRPr','endParaRPr']:return None
    attrs=[];ids=pkg.shape_map(part)
    for key,value in sorted(node.attrib.items()):
        name=E.QName(key).localname
        if active_format=='format-size' and name=='sz' and tag in ['rPr','defRPr','endParaRPr']:continue
        if (tag=='cNvPr' and name=='id') or name=='spid' or (tag in ['stCxn','endCxn'] and name=='id'):value='shape:'+ids.get(value,'MISSING:'+value)
        if name=='Requires' and tag=='Choice' or key=='{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable':
            value=json.dumps([(prefix,node.nsmap.get(prefix)) for prefix in value.split()])
        if key.startswith('{'+NS['r']+'}'):
            relation=pkg.rels.get(part,{}).get(value)
            if relation:value=json.dumps(relation,sort_keys=True)
        attrs.append((key,value))
    children=[value for child in node for value in [semantic(child,pkg,part,format_mode,format_target)] if value is not None]
    text=node.text or ''
    if not text.strip() and len(node):text=''
    # Applying a requested property may create pPr/defRPr and run-property
    # wrappers that did not exist before. Only this named target's formatting
    # mask removes wrappers left empty after removing the requested property;
    # any other paragraph/body/style attribute or field structure still counts.
    if active_format and tag in ['pPr','rPr','defRPr','endParaRPr'] and not attrs and not children and not text:return None
    return [node.tag,attrs,text,children]

def placeholder_chain(shape,pkg,part):
    ph=first(shape,'.//p:ph')
    if ph is None:return []
    chain=[];owner=part
    for kind in ['slideLayout','slideMaster']:
        owner=pkg.target(owner,kind)
        if not owner:break
        candidates=pkg.docs[owner].xpath('.//p:sp[p:nvSpPr/p:nvPr/p:ph]',namespaces=NS)
        match=next((s for s in candidates if first(s,'.//p:ph').get('idx','0')==ph.get('idx','0')),None)
        if match is None:match=next((s for s in candidates if first(s,'.//p:ph').get('type','body')==ph.get('type','body')),None)
        if match is not None:chain.append(match);ph=first(match,'.//p:ph')
    return chain

def geometries(shape,pkg,part,parent=(0.,0.,1.,1.),prefix=''):
    name=shape_name(shape);key=prefix+'/'+name
    transforms=[n for n in [transform_node(shape),*[transform_node(s) for s in placeholder_chain(shape,pkg,part)]] if n is not None]
    def item(tag):return next((n for t in transforms for n in t if local(n)==tag),None)
    def number(n,key,default=0):return float(n.get(key,default)) if n is not None else default
    off=item('off');ext=item('ext');x,y,sx,sy=parent
    rotation=next((float(t.get('rot')) for t in transforms if t.get('rot') is not None),0)
    result={key:[x+number(off,'x')*sx,y+number(off,'y')*sy,number(ext,'cx')*sx,number(ext,'cy')*sy,rotation, next((t.get('flipH') for t in transforms if t.get('flipH') is not None),'0'),next((t.get('flipV') for t in transforms if t.get('flipV') is not None),'0')]}
    if local(shape)=='grpSp':
        childoff=item('chOff');childext=item('chExt');gx=number(ext,'cx',1)/(number(childext,'cx',1) or 1);gy=number(ext,'cy',1)/(number(childext,'cy',1) or 1)
        transform=(x+(number(off,'x')-number(childoff,'x')*gx)*sx,y+(number(off,'y')-number(childoff,'y')*gy)*sy,sx*gx,sy*gy)
        for child in shape:
            if local(child) in SHAPES:result.update(geometries(child,pkg,part,transform,key))
    return result

def top_shapes(pkg,part):return [n for n in pkg.docs[part].xpath('./p:cSld/p:spTree/*',namespaces=NS) if local(n) in SHAPES]
def close_geometry(a,b,translate=False):
    if a.keys()!=b.keys():return False
    first_key=next(iter(a),None);dx=b[first_key][0]-a[first_key][0] if first_key and translate else 0;dy=b[first_key][1]-a[first_key][1] if first_key and translate else 0
    return all(all(abs((av+(dx if i==0 else dy if i==1 else 0))-bv)<=2 for i,(av,bv) in enumerate(zip(a[k][:5],b[k][:5]))) and a[k][5:]==b[k][5:] for k in a)

def raw_semantic(node):
    return [node.tag,sorted(node.attrib.items()),node.text or '',[raw_semantic(child) for child in node]] if node is not None else None

def audit(source,output,item,args):
    checks=[]
    def add(name,ok,detail=None):checks.append({'name':name,'passed':bool(ok),**({'detail':detail} if detail is not None else {})})
    add('source package integrity',not source.errors,source.errors);add('output package integrity',not output.errors,output.errors)
    add('independent python-pptx reader',output.readability['readable'],output.readability.get('error'))
    order=json.loads(args.slide_order) if args.slide_order else list(range(len(source.slides)))
    if args.scenario=='duplicate' and not args.slide_order:order=[0,0,*range(1,len(source.slides))]
    add('expected slide order length',len(output.slides)==len(order),{'expectedSourceIndices':order,'actualSlides':len(output.slides)})
    bytechecks=[]
    for part,data in source.parts.items():
        if any(fnmatch.fnmatchcase(part,pattern) for pattern in item.get('preserveParts',[])):
            same=output.parts.get(part)==data;bytechecks.append({'part':part,'sha256':hashlib.sha256(data).hexdigest(),'identical':same})
    add('untouched dependent parts preserve exact bytes',all(v['identical'] for v in bytechecks),bytechecks)
    for index,source_index in enumerate(order[:len(output.slides)]):
        if source_index>=len(source.slides):add('slide source index valid',False,source_index);continue
        sp=source.slides[source_index];op=output.slides[index];sd=source.docs[sp];od=output.docs[op]
        source_shapes=top_shapes(source,sp);output_shapes=top_shapes(output,op)
        shapechecks=[]
        for original in source_shapes:
            name=shape_name(original);matches=[n for n in output_shapes if shape_name(n)==name]
            if args.scenario in ['edit','duplicate'] and source_index==item.get('editPageIndex',0) and name==item.get('editObjectName','SAFE_EDIT'):
                if args.expected_edit_text:add(f'slide {index}: edited text',any(''.join(n.xpath('.//a:t/text()',namespaces=NS))==args.expected_edit_text for n in matches))
                continue
            mode=('format-color' if args.scenario=='nested-format-color' else args.scenario) if args.scenario in ['format-color','format-size','nested-format-color'] else None
            format_target=args.format_object_name or (item.get('codecNestedFormatObjectName') if args.scenario=='nested-format-color' else item.get('formatObjectName'))
            move=args.scenario=='move-object' and name==item.get('moveObjectName')
            eq=any(semantic(original,source,sp,mode,format_target)==semantic(n,output,op,mode,format_target) and close_geometry(geometries(original,source,sp),geometries(n,output,op),move) for n in matches)
            original_ids=[n.get('id') for n in original.xpath('.//p:cNvPr',namespaces=NS)]
            identity=any(close_geometry(geometries(original,source,sp),geometries(n,output,op),move) and [m.get('id') for m in n.xpath('.//p:cNvPr',namespaces=NS)]==original_ids for n in matches)
            shapechecks.append({'name':name,'preserved':eq and identity,'xmlAndGeometryPreserved':eq,'sourceShapeIdsRetained':identity,'matches':len(matches)})
        add(f'slide {index}: source shape XML semantics and resolved geometry',all(v['preserved'] for v in shapechecks),shapechecks)
        add(f'slide {index}: motion path, event/trigger, timing and build references',semantic(first(sd,'./p:timing'),source,sp)==semantic(first(od,'./p:timing'),output,op))
        if args.expected_edit_text and args.scenario in ['edit','duplicate'] and source_index==item.get('editPageIndex',0):
            add(f'slide {index}: requested edited text is present',args.expected_edit_text in ''.join(od.xpath('.//a:t/text()',namespaces=NS)))
        add(f'slide {index}: hidden state',sd.get('show','1')==od.get('show','1'))
        expected_notes=args.expected_notes if args.expected_notes is not None and source_index==item.get('editPageIndex',0) and args.scenario in ['edit','duplicate'] else notes_text(source,sp)
        add(f'slide {index}: notes',notes_text(output,op)==expected_notes)
        if args.scenario in ['format-color','format-size','nested-format-color'] and source_index==item.get('formatSlideIndex',0):
            name=args.format_object_name or (item.get('codecNestedFormatObjectName') if args.scenario=='nested-format-color' else item.get('formatObjectName'));target=next((n for n in od.iter() if local(n) in SHAPES and shape_name(n)==name),None)
            source_target=next((n for n in sd.iter() if local(n) in SHAPES and shape_name(n)==name),None)
            if target is not None and source_target is not None and any(local(n)=='grpSp' for n in source_target.iterancestors()):
                source_parent=next(n for n in source_target.iterancestors() if local(n)=='grpSp');output_parent=next((n for n in target.iterancestors() if local(n)=='grpSp'),None)
                add('nested format edit preserves every group and child raw transform',output_parent is not None and [raw_semantic(n) for n in source_parent.xpath('.//a:xfrm|.//p:xfrm',namespaces=NS)]==[raw_semantic(n) for n in output_parent.xpath('.//a:xfrm|.//p:xfrm',namespaces=NS)])
            runs=target.xpath('.//a:r|.//a:fld',namespaces=NS) if target is not None else []
            if args.scenario in ['format-color','nested-format-color']:
                expected=args.expected_color.lstrip('#').upper();values=[first(run,'./a:rPr/a:solidFill/a:srgbClr') for run in runs]
                add('requested color applied to each existing run and field',bool(values) and all(n is not None and n.get('val','').upper()==expected for n in values),expected)
            else:
                values=[first(run,'./a:rPr') for run in runs];expected=round(args.expected_font_size*100)
                add('requested font size applied to each existing run and field',bool(values) and all(n is not None and int(n.get('sz','0'))==expected for n in values),expected)
        if args.scenario=='move-object' and source_index==item.get('editPageIndex',0):
            name=item.get('moveObjectName');a=next((n for n in source_shapes if shape_name(n)==name),None);b=next((n for n in output_shapes if shape_name(n)==name),None)
            add('moved group exists',a is not None and b is not None)
            if a is not None and b is not None:
                ax=transform_node(a);bx=transform_node(b);ao=first(ax,'./a:off');bo=first(bx,'./a:off')
                add('moved group offset equals requested millimetres',ao is not None and bo is not None and abs(float(bo.get('x'))-float(ao.get('x'))-args.expected_dx*36000)<=2 and abs(float(bo.get('y'))-float(ao.get('y'))-args.expected_dy*36000)<=2)
                ag=copy.deepcopy(ax);bg=copy.deepcopy(bx)
                for node in [ag,bg]:
                    for off in node.xpath('./a:off',namespaces=NS):node.remove(off)
                add('moved group preserves rotation extents and child coordinate system',raw_semantic(ag)==raw_semantic(bg))
                ach=[n for n in a.xpath('.//a:xfrm|.//p:xfrm',namespaces=NS) if n is not ax];bch=[n for n in b.xpath('.//a:xfrm|.//p:xfrm',namespaces=NS) if n is not bx]
                add('moved group preserves every original child transform',list(map(raw_semantic,ach))==list(map(raw_semantic,bch)))
                add('moved group preserves complete original child XML', [raw_semantic(n) for n in a if local(n) in SHAPES]==[raw_semantic(n) for n in b if local(n) in SHAPES])
        if args.scenario=='duplicate-object' and source_index==item.get('editPageIndex',0):
            name=item.get('duplicateObjectName');original=next((n for n in source_shapes if shape_name(n)==name),None);copies=[n for n in output_shapes if shape_name(n)==name]
            add('duplicated object exists independently',original is not None and len(copies)==2)
            if original is not None:
                for n,group in enumerate(copies):
                    ids={m.get('id') for m in group.xpath('.//p:cNvPr',namespaces=NS)}
                    refs=group.xpath('.//a:stCxn/@id|.//a:endCxn/@id',namespaces=NS)
                    add(f'duplicated group {n}: connector targets its own children',all(target in ids for target in refs),{'localShapeIds':sorted(ids),'connectorTargets':refs})
                    add(f'duplicated group {n}: complete content and translated geometry',semantic(original,source,sp)==semantic(group,output,op) and close_geometry(geometries(original,source,sp),geometries(group,output,op),True))
    # Custom shows include duplicate occurrences; compare actual resolved slide targets, never just counts.
    for label,path in [('custom shows','./p:custShowLst'),('sections','./p:extLst')]:
        a=first(source.docs[source.presentation],path);b=first(output.docs[output.presentation],path)
        add(label+' and referenced slide destinations',semantic(a,source,source.presentation)==semantic(b,output,output.presentation))
    return {'scenario':args.scenario,'output':output.summary(),'checks':checks,'passed':all(c['passed'] for c in checks),'summary':{'checks':len(checks),'passed':sum(c['passed'] for c in checks),'failed':sum(not c['passed'] for c in checks)}}

def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('source');parser.add_argument('outputs',nargs='*');parser.add_argument('--fixture');parser.add_argument('--manifest',default=str(ROOT/'manifest.json'));parser.add_argument('--report');parser.add_argument('--scenario',choices=['noop','edit','duplicate','delete','format-color','format-size','duplicate-object','move-object','nested-format-color'],default='noop');parser.add_argument('--expected-edit-text');parser.add_argument('--expected-notes');parser.add_argument('--expected-color',default='#146CA4');parser.add_argument('--expected-font-size',type=float,default=28);parser.add_argument('--slide-order');parser.add_argument('--expected-dx',type=float,default=10);parser.add_argument('--expected-dy',type=float,default=0);parser.add_argument('--format-object-name')
    args=parser.parse_args();manifest=json.loads(Path(args.manifest).read_text());item=next((i for i in manifest['items'] if i['id']==args.fixture or Path(i['path']).name==Path(args.source).name),{})
    source=Package(args.source);result={'schemaVersion':1,'verifier':'Independent lxml + python-pptx; no production codec imports','fixture':item.get('id'),'source':source.summary(),'cycles':[]}
    if item.get('sha256') and source.sha256!=item['sha256']:source.errors.append({'kind':'source_digest_mismatch','expected':item['sha256'],'actual':source.sha256})
    for path in args.outputs:result['cycles'].append(audit(source,Package(path),item,args))
    result['passed']=not source.errors and source.readability['readable'] and all(c['passed'] for c in result['cycles'])
    if args.report:Path(args.report).parent.mkdir(parents=True,exist_ok=True);Path(args.report).write_text(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False)+'\n')
    print(json.dumps({'fixture':result['fixture'],'passed':result['passed'],'sourceErrors':source.errors,'outputs':[{'scenario':c['scenario'],**c['summary']} for c in result['cycles']]},ensure_ascii=False))
    for cycle in result['cycles']:
        for check in cycle['checks']:
            if not check['passed']:print(json.dumps({'failed':check['name'],'detail':check.get('detail')},ensure_ascii=False))
    return 0 if result['passed'] else 1
if __name__=='__main__':sys.exit(main())
