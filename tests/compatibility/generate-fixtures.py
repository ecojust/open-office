#!/usr/bin/env python3
"""Generate original compatibility samples; never label them as Office output.

SPDX-License-Identifier: MIT
Requires Python 3, python-docx, XlsxWriter, python-pptx and Pillow.
Only writes generated fixtures and their entries in this directory's manifest.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import struct
import wave
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo

from PIL import Image, ImageDraw
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import pptx
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PptInches, Pt as PptPt
import xlsxwriter
from lxml import etree

ROOT = Path(__file__).resolve().parent
GENERATED = ROOT / 'fixtures' / 'generated'
CREATED = datetime(2025, 1, 15, 12, 0, 0)
MANIFEST = ROOT / 'manifest.json'
ENTRIES = []


def normalize_zip(path):
    """Pin ZIP timestamps, without changing document XML or binary payloads."""
    with ZipFile(path) as source:
        entries = [(item.filename, source.read(item.filename)) for item in source.infolist()]
    with ZipFile(path, 'w') as output:
        for name, data in entries:
            info = ZipInfo(name, (2025, 1, 15, 12, 0, 0))
            info.compress_type = ZIP_STORED if name == 'mimetype' else ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            output.writestr(info, data)


def register(path, kind, features, assertions, generator):
    if path.suffix != '.rtf':
        normalize_zip(path)
    payload = path.read_bytes()
    ENTRIES.append({
        'id': path.stem,
        'path': str(path.relative_to(ROOT)),
        'kind': kind,
        'format': path.suffix[1:],
        'bytes': len(payload),
        'sha256': hashlib.sha256(payload).hexdigest(),
        'provenance': {
            'type': 'generated',
            'sourceUrl': None,
            'sourceCommit': None,
            'license': 'MIT',
            'author': 'Open Office project contributors',
            'generator': generator,
            'generatorPath': 'generate-fixtures.py',
            'producerDeclared': 'Generated test fixture; not a Microsoft Office, WPS or LibreOffice-authored original',
            'reproducibleTimestamp': '2025-01-15T12:00:00Z',
        },
        'features': features,
        'assertions': assertions,
    })


def brand_image():
    path = GENERATED / 'original-diagram.png'
    image = Image.new('RGB', (480, 180), '#eef4fb')
    draw = ImageDraw.Draw(image)
    for index, (label, color) in enumerate([('SOURCE', '#176d8f'), ('EDIT', '#7856a2'), ('SAVE', '#c06432')]):
        x = 14 + index * 160
        draw.rounded_rectangle((x, 38, x + 132, 141), 14, fill=color)
        draw.text((x + 21, 80), label, fill='white')
        if index < 2:
            draw.line((x + 135, 90, x + 151, 90), fill='#223344', width=4)
            draw.polygon([(x + 150, 85), (x + 157, 90), (x + 150, 95)], fill='#223344')
    image.save(path)
    return path


def set_doc_properties(document, title):
    document.core_properties.title = title
    document.core_properties.author = 'Open Office compatibility corpus generator'
    document.core_properties.subject = 'Synthetic original fixture, not Office-generated'
    document.core_properties.created = CREATED
    document.core_properties.modified = CREATED


def hyperlink(paragraph, label, url):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '176D8F'); props.append(color)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); paragraph._p.append(link)


def word_rich(image):
    document = docx.Document()
    set_doc_properties(document, 'COMPAT-WORD-RICH')
    section = document.sections[0]
    section.header.paragraphs[0].text = 'COMPAT-HEADER | 原创测试语料'
    footer = section.footer.paragraphs[0]
    footer.text = 'COMPAT-FOOTER | Page '
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    section.top_margin = Inches(.8); section.bottom_margin = Inches(.8)
    document.add_heading('COMPAT-WORD-RICH', 0)
    document.add_heading('江南项目 / Regional programme', 1)
    paragraph = document.add_paragraph('中文与 English coexist. ')
    run = paragraph.add_run('Bold blue original text'); run.bold = True; run.font.size = Pt(16); run.font.color.rgb = RGBColor.from_string('176D8F')
    document.add_paragraph('A numbered requirement', style='List Number')
    document.add_paragraph('A second numbered requirement', style='List Number')
    table = document.add_table(rows=4, cols=3); table.style = 'Light Shading Accent 1'
    table.cell(0, 0).merge(table.cell(0, 2)).text = 'MERGED-PROJECT-HEADER'
    for row, values in enumerate([['Region', 'Budget', 'Status'], ['江南', '12500.50', 'Approved'], ['North', '8400.00', 'Review']], 1):
        for col, value in enumerate(values): table.cell(row, col).text = value
    document.add_picture(str(image), width=Inches(4.8))
    hyperlink(document.add_paragraph('Reference: '), 'Offline-safe example hyperlink', 'https://example.invalid/compatibility-reference')
    document.add_page_break()
    document.add_heading('APPENDIX-PERSISTENT', 1)
    document.add_paragraph('This appendix must remain editable after repeated save and reopen cycles.')
    path = GENERATED / 'writer-rich.docx'; document.save(path)
    register(path, 'word', ['bilingual', 'heading-styles', 'numbered-list', 'merged-table', 'header-footer', 'page-number-field', 'image', 'external-hyperlink-no-fetch', 'page-break'], {
        'textIncludes': ['COMPAT-WORD-RICH', '江南项目', 'MERGED-PROJECT-HEADER', 'APPENDIX-PERSISTENT'],
        'minimumFeatures': {'tables': 1, 'images': 1, 'hyperlinks': 1},
    }, f'python-docx {docx.__version__}; original Pillow diagram')


def workbook_rich(image):
    path = GENERATED / 'calc-rich.xlsx'
    book = xlsxwriter.Workbook(path)
    book.set_properties({'title': 'COMPAT-CALC-RICH', 'author': 'Open Office compatibility corpus generator', 'created': CREATED})
    title = book.add_format({'bold': True, 'font_size': 17, 'font_color': '#176D8F'})
    money = book.add_format({'num_format': '#,##0.00', 'font_color': '#176D8F'})
    head = book.add_format({'bold': True, 'bg_color': '#DBEAF5', 'border': 1})
    summary = book.add_worksheet('Summary'); data = book.add_worksheet('Data'); archive = book.add_worksheet('Archive')
    summary.merge_range('A1:D1', 'COMPAT-CALC-RICH', title)
    summary.write('A3', '江南项目'); summary.write_formula('B3', '=SUM(Data!D2:D5)', money, 228)
    summary.insert_image('F2', str(image), {'x_scale': .45, 'y_scale': .45})
    summary.write_url('A5', 'https://example.invalid/report', string='Reference without network loading')
    data.write_row('A1', ['Region', 'Quantity', 'Unit price', 'Revenue'], head)
    rows = [('江南', 2, 40), ('North', 3, 12), ('江南', 4, 10), ('West', 6, 12)]
    for index, (region, quantity, price) in enumerate(rows, 1):
        data.write_row(index, 0, [region, quantity, price])
        data.write_formula(index, 3, f'=B{index+1}*C{index+1}', money, quantity * price)
    data.freeze_panes(1, 1); data.autofilter('A1:D5'); data.set_column('A:D', 18)
    data.data_validation('B2:B30', {'validate': 'integer', 'criteria': 'between', 'minimum': 0, 'maximum': 999})
    data.conditional_format('D2:D5', {'type': '3_color_scale'})
    data.write_comment('D2', 'A real worksheet comment / 原生批注')
    data.write_string('F2', '00123'); data.write_string('F3', '=SUM(9,9)')
    archive.write('A1', 'HIDDEN-ARCHIVE-SENTINEL'); archive.hide()
    book.define_name('RevenueRange', '=Data!$D$2:$D$5')
    chart = book.add_chart({'type': 'column'})
    chart.add_series({'name': 'Revenue', 'categories': '=Data!$A$2:$A$5', 'values': '=Data!$D$2:$D$5'})
    chart.set_title({'name': 'Regional revenue'}); summary.insert_chart('A8', chart)
    book.close()
    register(path, 'sheet', ['bilingual', 'three-sheets', 'hidden-sheet', 'formula-cache', 'cross-sheet-formula', 'merged-cells', 'number-format', 'data-validation', 'conditional-format', 'chart', 'named-range', 'comment', 'image', 'hyperlink', 'freeze-panes', 'autofilter', 'literal-formula-text'], {
        'sheetNames': ['Summary', 'Data', 'Archive'],
        'sheetCells': [
            {'sheet': 'Summary', 'address': 'A1', 'type': 'text', 'value': 'COMPAT-CALC-RICH'},
            {'sheet': 'Summary', 'address': 'B3', 'type': 'formula', 'formula': 'SUM(Data!D2:D5)', 'value': 228},
            {'sheet': 'Data', 'address': 'D2', 'type': 'formula', 'formula': 'B2*C2', 'value': 80},
            {'sheet': 'Data', 'address': 'F2', 'type': 'text', 'value': '00123'},
            {'sheet': 'Data', 'address': 'F3', 'type': 'text', 'value': '=SUM(9,9)'},
        ],
        'minimumFeatures': {'charts': 1, 'images': 1, 'hyperlinks': 1},
    }, f'XlsxWriter {xlsxwriter.__version__}; original Pillow diagram')


def presentation_rich(image):
    presentation = Presentation()
    set_doc_properties(presentation, 'COMPAT-PRESENTATION-RICH')
    presentation.slide_width = PptInches(13.333); presentation.slide_height = PptInches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = 'COMPAT-PRESENTATION-RICH'
    slide.placeholders[1].text = '江南项目 / An editable original presentation'
    slide.notes_slide.notes_text_frame.text = 'SPEAKER-NOTE-ONE / 保留这条备注'
    slide = presentation.slides.add_slide(presentation.slide_layouts[5]); slide.shapes.title.text = 'Editable data / 可编辑数据'
    table = slide.shapes.add_table(4, 3, PptInches(.5), PptInches(1.5), PptInches(5), PptInches(2.4)).table
    for row, values in enumerate([['Region', '2024', '2025'], ['江南', '24', '30'], ['North', '18', '22'], ['West', '15', '21']]):
        for col, value in enumerate(values): table.cell(row, col).text = value
    data = CategoryChartData(); data.categories = ['江南', 'North', 'West']; data.add_series('2024', [24, 18, 15]); data.add_series('2025', [30, 22, 21])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, PptInches(6), PptInches(1.4), PptInches(6.5), PptInches(4.8), data)
    slide.notes_slide.notes_text_frame.text = 'SPEAKER-NOTE-TWO / Table and chart remain objects.'
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(PptInches(.6), PptInches(.5), PptInches(11), PptInches(1))
    box.text_frame.paragraphs[0].text = 'PROCESS-THREE-STEPS'
    box.text_frame.paragraphs[0].runs[0].font.size = PptPt(28)
    slide.shapes.add_picture(str(image), PptInches(.7), PptInches(2), width=PptInches(8))
    link = slide.shapes.add_textbox(PptInches(9), PptInches(2), PptInches(3), PptInches(1.5)).text_frame.paragraphs[0].add_run()
    link.text = 'Offline-safe hyperlink'; link.hyperlink.address = 'https://example.invalid/presentation'
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(9), PptInches(4), PptInches(3), PptInches(1))
    shape.text = 'Editable shape'
    slide.notes_slide.notes_text_frame.text = 'SPEAKER-NOTE-THREE'
    path = GENERATED / 'presentation-rich.pptx'; presentation.save(path)
    register(path, 'presentation', ['bilingual', 'three-layouts', 'editable-table', 'editable-chart', 'image', 'hyperlink', 'shape', 'speaker-notes'], {
        'textIncludes': ['COMPAT-PRESENTATION-RICH', 'Editable data', 'PROCESS-THREE-STEPS'],
        'slideCount': 3,
        'minimumFeatures': {'tables': 1, 'charts': 1, 'images': 1, 'hyperlinks': 1},
    }, f'python-pptx {pptx.__version__}; original Pillow diagram')



def word_long(image):
    document = docx.Document()
    set_doc_properties(document, 'COMPAT-LONG-REPORT')
    for index in range(20):
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        section.orientation = WD_ORIENT.LANDSCAPE if index % 4 == 3 else WD_ORIENT.PORTRAIT
        section.page_width = Inches(11.7 if index % 4 == 3 else 8.3)
        section.page_height = Inches(8.3 if index % 4 == 3 else 11.7)
        section.top_margin = section.bottom_margin = Inches(.65)
        section.header.is_linked_to_previous = False
        section.header.paragraphs[0].text = f'COMPAT-LONG-REPORT / SECTION-{index + 1:02d}'
        document.add_heading(f'LONG-PAGE-{index + 1:02d}', 1)
        document.add_paragraph('江南地区季度报告 / A multilingual section with independently editable figures.')
        for paragraph in range(3):
            document.add_paragraph(f'Finding {paragraph + 1}: preserve this section, its table and its page orientation after editing. 字体、分页与页边距均属于测试对象。')
        table = document.add_table(rows=5, cols=4); table.style = 'Table Grid'
        for row in range(5):
            for col in range(4): table.cell(row, col).text = f'S{index+1:02d}-R{row+1}-C{col+1}'
        if index % 4 == 0: document.add_picture(str(image), width=Inches(3.6))
    path = GENERATED / 'writer-long-report.docx'; document.save(path)
    register(path, 'word', ['20-explicit-page-sections', 'mixed-portrait-landscape', '20-tables', 'section-specific-headers', 'repeated-image', 'bilingual'], {
        'textIncludes': ['LONG-PAGE-01', 'LONG-PAGE-10', 'LONG-PAGE-20', 'S20-R5-C4'],
        'minimumFeatures': {'tables': 20, 'images': 1},
    }, f'python-docx {docx.__version__}; 20 compact explicit sections, rendered page count must be measured')


def rewrite_package(path, changes):
    with ZipFile(path) as source: contents = {name: source.read(name) for name in source.namelist()}
    contents.update(changes)
    with ZipFile(path, 'w', ZIP_DEFLATED) as output:
        for name, data in contents.items(): output.writestr(name, data)


def xml_bytes(node):
    return etree.tostring(node, encoding='UTF-8', xml_declaration=True, standalone=True)


def word_review(image):
    document = docx.Document(); set_doc_properties(document, 'COMPAT-REVIEW-DOCX')
    document.add_heading('COMPAT-REVIEW-DOCX', 0)
    paragraph = document.add_paragraph('Review this proposal: ')
    run = paragraph.add_run('COMMENT-ANCHOR')
    comment = document.add_comment(run, text='EDITOR-COMMENT / 请核对这项数据。', author='Compatibility Author', initials='QA')
    comment._element.set(qn('w:date'), '2025-01-15T12:00:00Z')
    paragraph = document.add_paragraph('Revision sample: ')
    for index, (kind, text) in enumerate([('del', 'OLD-REVISED-TEXT'), ('ins', 'NEW-REVISED-TEXT')]):
        change = OxmlElement(f'w:{kind}'); change.set(qn('w:id'), str(20+index)); change.set(qn('w:author'), 'Fixture Author'); change.set(qn('w:date'), '2025-01-15T12:00:00Z')
        run = OxmlElement('w:r'); value = OxmlElement('w:delText' if kind == 'del' else 'w:t'); value.text = text; run.append(value); change.append(run); paragraph._p.append(change)
    paragraph = document.add_paragraph('Reference with footnote')
    run = OxmlElement('w:r'); ref = OxmlElement('w:footnoteReference'); ref.set(qn('w:id'), '1'); run.append(ref); paragraph._p.append(run)
    outer = document.add_table(rows=2, cols=2); outer.style='Table Grid'; outer.cell(0,0).merge(outer.cell(0,1)).text='COMPLEX-MERGED-HEADER'
    nested = outer.cell(1,0).add_table(rows=2, cols=2)
    for i in range(2):
        for j in range(2): nested.cell(i,j).text=f'NESTED-{i}-{j}'
    outer.cell(1,1).text='OUTER-CELL-PERSISTENT'
    document.add_picture(str(image), width=Inches(3))
    path=GENERATED/'writer-review-footnotes.docx'; document.save(path)
    w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    notes=etree.Element('{'+w+'}footnotes',nsmap={'w':w})
    for ident, separator in [(-1,'separator'),(0,'continuationSeparator')]:
        note=etree.SubElement(notes,'{'+w+'}footnote',{'{'+w+'}id':str(ident),'{'+w+'}type':separator})
        etree.SubElement(etree.SubElement(etree.SubElement(note,'{'+w+'}p'),'{'+w+'}r'),'{'+w+'}'+separator)
    note=etree.SubElement(notes,'{'+w+'}footnote',{'{'+w+'}id':'1'})
    etree.SubElement(etree.SubElement(etree.SubElement(note,'{'+w+'}p'),'{'+w+'}r'),'{'+w+'}t').text='FOOTNOTE-PERSISTENT / 原创脚注'
    with ZipFile(path) as source:
        rel=etree.fromstring(source.read('word/_rels/document.xml.rels')); types=etree.fromstring(source.read('[Content_Types].xml'))
    etree.SubElement(rel,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='rIdCompatFootnote',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes',Target='footnotes.xml')
    etree.SubElement(types,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',PartName='/word/footnotes.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    rewrite_package(path,{'word/footnotes.xml':xml_bytes(notes),'word/_rels/document.xml.rels':xml_bytes(rel),'[Content_Types].xml':xml_bytes(types)})
    register(path,'word',['bilingual','comment','tracked-insertion','tracked-deletion','footnote','nested-table','merged-table','image'],{
        'textIncludes':['COMPAT-REVIEW-DOCX','COMMENT-ANCHOR','OUTER-CELL-PERSISTENT'],
        'minimumFeatures':{'tables':2,'images':1,'comments':1,'footnotes':1,'trackedChanges':2},
    }, f'python-docx {docx.__version__}; explicit original OOXML review and footnote parts')


def workbook_large(image):
    path=GENERATED/'calc-2000-rows.xlsx'; book=xlsxwriter.Workbook(path)
    book.set_properties({'title':'COMPAT-LARGE-CALC','author':'Open Office compatibility corpus generator','created':CREATED})
    dash=book.add_worksheet('Dashboard'); data=book.add_worksheet('Transactions'); tax=book.add_worksheet('TaxRates'); hidden=book.add_worksheet('HiddenCalc')
    money=book.add_format({'num_format':'#,##0.00'}); title=book.add_format({'bold':True,'font_size':18})
    dash.write('A1','COMPAT-LARGE-CALC',title); tax.write_row('A1',['Jurisdiction','Rate']); tax.write_row('A2',['Local',.06]); hidden.write('A1','LARGE-HIDDEN-SENTINEL'); hidden.hide()
    data.add_table('A1:F2001',{'name':'SalesLedger','columns':[{'header':name} for name in ['ID','Region','Quantity','Price','Gross','WithTax']]})
    gross=0
    for row in range(1,2001):
        quantity=row%9+1; price=10+row%17; amount=quantity*price; gross+=amount
        data.write_row(row,0,[f'TXN-{row:05d}',['江南','North','West'][row%3],quantity,price])
        data.write_formula(row,4,f'=C{row+1}*D{row+1}',money,amount)
        data.write_formula(row,5,f'=E{row+1}*(1+TaxRates!$B$2)',money,amount*1.06)
    dash.write('A3','Gross total'); dash.write_formula('B3','=SUM(Transactions!E2:E2001)',money,gross)
    dash.write_formula('B4','=COUNTA(Transactions!A2:A2001)',None,2000)
    dash.write('A7','First ten transactions')
    chart=book.add_chart({'type':'line'}); chart.add_series({'name':'Gross','categories':'=Transactions!$A$2:$A$11','values':'=Transactions!$E$2:$E$11'}); dash.insert_chart('A8',chart)
    book.define_name('GrossTransactions','=Transactions!$E$2:$E$2001')
    data.freeze_panes(1,2); data.set_column('A:B',18); data.set_column('C:F',14)
    data.conditional_format('E2:E2001',{'type':'data_bar'}); data.data_validation('C2:C2001',{'validate':'integer','criteria':'between','minimum':1,'maximum':9})
    data.repeat_rows(0); data.fit_to_pages(1,0); data.print_area('A1:F2001'); book.close()
    register(path,'sheet',['2000-data-rows','4000-row-formulas','four-sheets','cross-sheet-formulas','table','line-chart','hidden-sheet','named-range','validation','conditional-format','repeat-print-header'],{
        'sheetNames':['Dashboard','Transactions','TaxRates','HiddenCalc'],
        'sheetCells':[{'sheet':'Dashboard','address':'A1','type':'text','value':'COMPAT-LARGE-CALC'},{'sheet':'Dashboard','address':'B3','type':'formula','value':gross},{'sheet':'Dashboard','address':'B4','type':'formula','value':2000},{'sheet':'Transactions','address':'A2001','type':'text','value':'TXN-02000'}],
        'minimumFeatures':{'charts':1},
    },f'XlsxWriter {xlsxwriter.__version__}; deterministic 2000-row ledger')


def workbook_dates(image):
    path=GENERATED/'calc-dates-validation.xlsx'; book=xlsxwriter.Workbook(path,{'date_1904':True})
    book.set_properties({'title':'COMPAT-DATE-SYSTEM','author':'Open Office compatibility corpus generator','created':CREATED})
    sheet=book.add_worksheet('Calendar'); lookup=book.add_worksheet('Lists'); lookup.write_column('A1',['待办','进行中','完成'])
    sheet.write('A1','COMPAT-DATE-SYSTEM'); sheet.write('A3','Date stored using 1904 epoch'); sheet.write_datetime('B3',datetime(2024,2,29,12,30),book.add_format({'num_format':'yyyy-mm-dd hh:mm'}))
    sheet.write_string('A5','0000123'); sheet.write_string('A6','=1+1'); sheet.write_boolean('B5',True); sheet.write_formula('B6','=IF(B5,42,0)',None,42)
    sheet.write_row('A8',['Task','Status','Progress']); sheet.write_row('A9',['中文任务','进行中',.5],book.add_format({'text_wrap':True})); sheet.write_row('A10',['Review','完成',1])
    sheet.data_validation('B9:B25',{'validate':'list','source':'=StateOptions'}); book.define_name('StateOptions','=Lists!$A$1:$A$3')
    sheet.conditional_format('C9:C25',{'type':'icon_set','icon_style':'3_traffic_lights'}); sheet.set_column('C:C',15,book.add_format({'num_format':'0%'}))
    sheet.set_row(12,None,None,{'hidden':True}); sheet.write('A13','HIDDEN-ROW-SENTINEL'); sheet.set_column('F:F',12,None,{'hidden':True}); sheet.write('F4','HIDDEN-COLUMN-SENTINEL')
    sheet.write_comment('B3','Leap day / 1904 epoch test'); lookup.very_hidden(); book.close()
    register(path,'sheet',['1904-date-system','leap-day','date-number-format','boolean','formula','literal-leading-zero','literal-equals-text','named-validation-list','icon-conditional-format','hidden-row-column','very-hidden-sheet','comment'],{
        'sheetNames':['Calendar','Lists'],
        'sheetCells':[{'sheet':'Calendar','address':'A1','type':'text','value':'COMPAT-DATE-SYSTEM'},{'sheet':'Calendar','address':'A5','type':'text','value':'0000123'},{'sheet':'Calendar','address':'A6','type':'text','value':'=1+1'},{'sheet':'Calendar','address':'B6','type':'formula','value':42}],
    },f'XlsxWriter {xlsxwriter.__version__}; 1904 date epoch and validation edge cases')


def presentation_long(image):
    presentation=Presentation(); set_doc_properties(presentation,'COMPAT-TWENTY-SLIDES')
    presentation.slide_width=PptInches(13.333); presentation.slide_height=PptInches(7.5)
    for index in range(20):
        slide=presentation.slides.add_slide(presentation.slide_layouts[[0,1,5,6][index%4]])
        title=f'DECK-PAGE-{index+1:02d}'
        if slide.shapes.title is not None: slide.shapes.title.text=title
        else: slide.shapes.add_textbox(PptInches(.5),PptInches(.4),PptInches(11),PptInches(.8)).text=title
        box=slide.shapes.add_textbox(PptInches(.5),PptInches(1.4),PptInches(11),PptInches(.7)); box.text='江南计划 / Editable original briefing'
        if index%4==0:
            table=slide.shapes.add_table(3,3,PptInches(.7),PptInches(2.5),PptInches(5),PptInches(2.2)).table
            for row in range(3):
                for col in range(3): table.cell(row,col).text=f'P{index+1}-R{row+1}-C{col+1}'
        elif index%4==1:
            data=CategoryChartData(); data.categories=['A','B','C']; data.add_series('Values',[index+1,index+3,index+6])
            slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,PptInches(.7),PptInches(2.4),PptInches(8),PptInches(4),data)
        else: slide.shapes.add_picture(str(image),PptInches(1),PptInches(3),width=PptInches(8))
        if index in (5,11): slide._element.set('show','0')
        slide.notes_slide.notes_text_frame.text=f'LONG-SPEAKER-NOTE-{index+1:02d} / 本页讲解要点'
    path=GENERATED/'presentation-20-slides.pptx'; presentation.save(path)
    register(path,'presentation',['20-slides','four-layouts','5-editable-tables','5-editable-charts','10-images','20-speaker-notes','2-hidden-slides','bilingual'],{
        'textIncludes':['DECK-PAGE-01','DECK-PAGE-10','DECK-PAGE-20'],'slideCount':20,'minimumFeatures':{'tables':5,'charts':5,'images':1},
    },f'python-pptx {pptx.__version__}; original 20-slide briefing')


def presentation_effects(image):
    presentation=Presentation(); set_doc_properties(presentation,'COMPAT-EFFECTS-MEDIA')
    slide=presentation.slides.add_slide(presentation.slide_layouts[5]); slide.shapes.title.text='COMPAT-EFFECTS-MEDIA'
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PptInches(1),PptInches(2),PptInches(5),PptInches(1.2)); shape.text='ANIMATED-EDITABLE-SHAPE'
    slide.notes_slide.notes_text_frame.text='This slide has an explicit fade entrance and a transition sound; playback is a separate compatibility capability.'
    second=presentation.slides.add_slide(presentation.slide_layouts[5]); second.shapes.title.text='SECOND-TIMED-SLIDE'; second.shapes.add_picture(str(image),PptInches(1),PptInches(2),width=PptInches(6))
    path=GENERATED/'presentation-effects-media.pptx'; presentation.save(path)
    audio=io.BytesIO()
    with wave.open(audio,'wb') as output:
        output.setnchannels(1); output.setsampwidth(2); output.setframerate(8000)
        output.writeframes(b''.join(struct.pack('<h',int(2500*math.sin(2*math.pi*440*i/8000)*min(1,i/400,(8000-i)/400))) for i in range(8000)))
    p='http://schemas.openxmlformats.org/presentationml/2006/main'; r='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with ZipFile(path) as source:
        xml=etree.fromstring(source.read('ppt/slides/slide1.xml')); rel=etree.fromstring(source.read('ppt/slides/_rels/slide1.xml.rels')); types=etree.fromstring(source.read('[Content_Types].xml'))
    transition=etree.fromstring(f'<p:transition xmlns:p="{p}" xmlns:r="{r}" spd="slow" advClick="1" advTm="1500"><p:fade/><p:sndAc><p:stSnd><p:snd r:embed="rIdCompatAudio" name="compat-tone.wav"/></p:stSnd></p:sndAc></p:transition>')
    timing=etree.fromstring(f'<p:timing xmlns:p="{p}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="4" dur="1000" fill="hold"/><p:tgtEl><p:spTgt spid="{shape.shape_id}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    xml.append(transition); xml.append(timing)
    etree.SubElement(rel,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='rIdCompatAudio',Type=r+'/audio',Target='../media/compat-tone.wav')
    etree.SubElement(types,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='wav',ContentType='audio/wav')
    rewrite_package(path,{'ppt/slides/slide1.xml':xml_bytes(xml),'ppt/slides/_rels/slide1.xml.rels':xml_bytes(rel),'[Content_Types].xml':xml_bytes(types),'ppt/media/compat-tone.wav':audio.getvalue()})
    register(path,'presentation',['editable-shape','fade-entrance-animation','slide-transition','timed-advance','embedded-original-wav-transition-sound','speaker-notes','image'],{
        'textIncludes':['COMPAT-EFFECTS-MEDIA','ANIMATED-EDITABLE-SHAPE','SECOND-TIMED-SLIDE'],'slideCount':2,'minimumFeatures':{'images':1},
    },f'python-pptx {pptx.__version__}; explicit original OOXML timing and 1-second synthesized WAV')


def rtf_text(value):
    return ''.join('\\'+c if c in '\\{}' else c if ord(c)<128 else f'\\u{ord(c) if ord(c)<32768 else ord(c)-65536}?' for c in value)


def word_rtf(image):
    contents=r'{\rtf1\ansi\ansicpg1252\deff0\uc1{\fonttbl{\f0 Arial;}{\f1 Noto Serif CJK SC;}}{\colortbl;\red23\green109\blue143;}'
    contents+=r'{\header\pard RTF-HEADER-PERSISTENT\par}{\footer\pard RTF-FOOTER-PERSISTENT\par}'
    contents+=r'\paperw11906\paperh16838\margl1440\margr1440\pard\b\fs36 COMPAT-RICH-RTF\b0\fs24\par '
    contents+=rtf_text('江南项目 / Original Unicode RTF sample.')+r'\par\cf1 A blue sentence.\cf0\par '
    for row in [('Item','Count'),('RTF-TABLE-SENTINEL','42')]:
        contents+=r'\trowd\trgaph108\cellx3600\cellx7200\intbl '+rtf_text(row[0])+r'\cell '+rtf_text(row[1])+r'\cell\row '
    contents+=r'\pard{\pict\pngblip\picw480\pich180\picwgoal4800\pichgoal1800 '+image.read_bytes().hex()+r'}\par\page SECOND-RTF-PAGE\par}'
    path=GENERATED/'writer-unicode-tables.rtf'; path.write_text(contents,encoding='ascii')
    register(path,'word',['rtf-unicode','table','header-footer','image','text-format','explicit-page-break'],{'textIncludes':['COMPAT-RICH-RTF','RTF-TABLE-SENTINEL','SECOND-RTF-PAGE'],'minimumFeatures':{'images':1}},'Original standards-based RTF generator with Pillow PNG')


ODF_NS={
    'office':'urn:oasis:names:tc:opendocument:xmlns:office:1.0','text':'urn:oasis:names:tc:opendocument:xmlns:text:1.0','table':'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
    'draw':'urn:oasis:names:tc:opendocument:xmlns:drawing:1.0','style':'urn:oasis:names:tc:opendocument:xmlns:style:1.0','fo':'urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0',
    'svg':'urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0','xlink':'http://www.w3.org/1999/xlink','presentation':'urn:oasis:names:tc:opendocument:xmlns:presentation:1.0',
    'meta':'urn:oasis:names:tc:opendocument:xmlns:meta:1.0','dc':'http://purl.org/dc/elements/1.1/','of':'urn:oasis:names:tc:opendocument:xmlns:of:1.2'}
ODF_DECL=' '.join(f'xmlns:{key}="{value}"' for key,value in ODF_NS.items())


def write_odf(path, kind, body, automatic, image=None):
    mime='application/vnd.oasis.opendocument.'+{'odt':'text','ods':'spreadsheet','odp':'presentation'}[kind]
    content=f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {ODF_DECL} office:version="1.2"><office:scripts/><office:font-face-decls><style:font-face style:name="CompatSans" svg:font-family="Arial"/></office:font-face-decls><office:automatic-styles>{automatic}</office:automatic-styles><office:body>{body}</office:body></office:document-content>'
    page=f'<style:page-layout style:name="Page"><style:page-layout-properties fo:page-width="{28 if kind=="odp" else 21}cm" fo:page-height="{15.75 if kind=="odp" else 29.7}cm" style:print-orientation="{"landscape" if kind=="odp" else "portrait"}" fo:margin="1.5cm"/></style:page-layout>'
    styles=f'<?xml version="1.0" encoding="UTF-8"?><office:document-styles {ODF_DECL} office:version="1.2"><office:styles><style:style style:name="Standard" style:family="paragraph"><style:text-properties style:font-name="CompatSans" fo:font-size="12pt"/></style:style><style:style style:name="Title" style:family="paragraph" style:parent-style-name="Standard"><style:text-properties fo:font-size="22pt" fo:font-weight="bold" fo:color="#176d8f"/></style:style></office:styles><office:automatic-styles>{page}</office:automatic-styles><office:master-styles><style:master-page style:name="Default" style:page-layout-name="Page">'+('<style:header><text:p>ODF-HEADER-PERSISTENT</text:p></style:header><style:footer><text:p>ODF-FOOTER-PERSISTENT</text:p></style:footer>' if kind=='odt' else '')+'</style:master-page></office:master-styles></office:document-styles>'
    meta=f'<office:document-meta {ODF_DECL} office:version="1.2"><office:meta><meta:generator>Open Office original fixture generator (not LibreOffice)</meta:generator><dc:creator>Open Office compatibility corpus generator</dc:creator><meta:creation-date>2025-01-15T12:00:00</meta:creation-date></office:meta></office:document-meta>'
    parts={'mimetype':mime.encode(),'content.xml':content.encode(),'styles.xml':styles.encode(),'meta.xml':meta.encode()}
    if image: parts['Pictures/original-diagram.png']=image.read_bytes()
    manifest='<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="'+mime+'" manifest:version="1.2"/>'
    for name in parts:
        if name!='mimetype': manifest+=f'<manifest:file-entry manifest:full-path="{name}" manifest:media-type="{"image/png" if name.endswith("png") else "text/xml"}"/>'
    parts['META-INF/manifest.xml']=(manifest+'</manifest:manifest>').encode()
    with ZipFile(path,'w') as output:
        for name,data in parts.items(): output.writestr(name,data,compress_type=ZIP_STORED if name=='mimetype' else ZIP_DEFLATED)


def odf_image(x='1cm',y='5cm',anchor='paragraph'):
    return f'<draw:frame draw:name="OriginalDiagram" text:anchor-type="{anchor}" svg:x="{x}" svg:y="{y}" svg:width="12cm" svg:height="4.5cm"><draw:image xlink:href="Pictures/original-diagram.png" xlink:type="simple" xlink:show="embed" xlink:actuate="onLoad"/></draw:frame>'


def odf_table(name,rows):
    result=f'<table:table table:name="{name}"><table:table-column table:number-columns-repeated="{len(rows[0])}"/>'
    for row in rows:
        result+='<table:table-row>'
        for value in row: result+='<table:table-cell office:value-type="string"><text:p>'+escape(str(value))+'</text:p></table:table-cell>'
        result+='</table:table-row>'
    return result+'</table:table>'


def word_odf(image):
    body='<office:text><text:h text:outline-level="1" text:style-name="Title">COMPAT-ORIGINAL-ODT</text:h><text:p>江南项目 / Native ODF structure, generated from the public format specification.</text:p>'
    body+=odf_table('Project',[['Area','Budget'],['江南','4200'],['ODT-TABLE-SENTINEL','7600']])
    body+='<text:p>With an original illustration '+odf_image('0cm','0cm','as-char')+'</text:p><text:p><text:a xlink:href="https://example.invalid/odt" xlink:type="simple">Offline-safe ODF hyperlink</text:a></text:p>'
    body+='<text:p>ODT note<text:note text:id="note1" text:note-class="footnote"><text:note-citation>1</text:note-citation><text:note-body><text:p>ODT-FOOTNOTE-PERSISTENT</text:p></text:note-body></text:note></text:p><text:p text:style-name="Break">ODT-SECOND-PAGE</text:p></office:text>'
    path=GENERATED/'writer-original-odf.odt'; write_odf(path,'odt',body,'<style:style style:name="Break" style:family="paragraph"><style:paragraph-properties fo:break-before="page"/></style:style>',image)
    register(path,'word',['original-odf','bilingual','table','image','hyperlink','header-footer','footnote','explicit-page-break'],{'textIncludes':['COMPAT-ORIGINAL-ODT','ODT-TABLE-SENTINEL','ODT-SECOND-PAGE'],'minimumFeatures':{'tables':1,'images':1,'footnotes':1,'hyperlinks':1}},'Original ODF 1.2 XML/ZIP generator; not LibreOffice-authored')


def workbook_odf(image):
    body='<office:spreadsheet>'+odf_table('Budget',[['COMPAT-ORIGINAL-ODS','Amount'],['江南','1250'],['North','900']])
    body+='<table:table table:name="Calculations"><table:table-column table:number-columns-repeated="3"/><table:table-row><table:table-cell office:value-type="string"><text:p>Total</text:p></table:table-cell><table:table-cell office:value-type="float" office:value="2150" table:formula="of:=SUM([Budget.B2:Budget.B3])"><text:p>2150</text:p></table:table-cell></table:table-row><table:table-row><table:table-cell office:value-type="string"><text:p>Literal identifier</text:p></table:table-cell><table:table-cell office:value-type="string"><text:p>000987</text:p></table:table-cell></table:table-row></table:table>'
    # Numeric values, rather than numeric-looking strings, are needed for the SUM source.
    body=body.replace('<table:table-cell office:value-type="string"><text:p>1250</text:p>','<table:table-cell office:value-type="float" office:value="1250"><text:p>1250</text:p>').replace('<table:table-cell office:value-type="string"><text:p>900</text:p>','<table:table-cell office:value-type="float" office:value="900"><text:p>900</text:p>')
    body+='<table:named-expressions><table:named-range table:name="RegionalBudget" table:base-cell-address="$Budget.$B$2" table:cell-range-address="$Budget.$B$2:.$B$3"/></table:named-expressions></office:spreadsheet>'
    path=GENERATED/'calc-original-odf.ods'; write_odf(path,'ods',body,'')
    register(path,'sheet',['original-odf','two-sheets','cross-sheet-formula','numeric-cached-result','named-range','bilingual','literal-leading-zero'],{'sheetNames':['Budget','Calculations'],'sheetCells':[{'sheet':'Budget','address':'A1','type':'text','value':'COMPAT-ORIGINAL-ODS'},{'sheet':'Calculations','address':'B1','type':'formula','value':2150},{'sheet':'Calculations','address':'B2','type':'text','value':'000987'}]},'Original ODF 1.2 XML/ZIP generator; not LibreOffice-authored')


def presentation_odf(image):
    body='<office:presentation>'
    for index in range(3):
        title=['COMPAT-ORIGINAL-ODP','ODP-DATA-SLIDE','ODP-FINAL-SLIDE'][index]
        body+=f'<draw:page draw:name="Page{index+1}" draw:master-page-name="Default"><draw:frame draw:name="Title{index+1}" svg:x="1cm" svg:y="1cm" svg:width="24cm" svg:height="2cm"><draw:text-box><text:p text:style-name="Title">{title}</text:p></draw:text-box></draw:frame>'
        if index==0: body+=odf_image('1cm','4cm')
        elif index==1: body+='<draw:frame draw:name="DataTable" svg:x="1cm" svg:y="4cm" svg:width="16cm" svg:height="4cm">'+odf_table('SlideData',[['Region','Amount'],['江南','42'],['North','31']])+'</draw:frame>'
        else: body+='<draw:rect draw:name="EditableRectangle" svg:x="2cm" svg:y="5cm" svg:width="8cm" svg:height="3cm"><text:p>EDITABLE-ODP-SHAPE</text:p></draw:rect><draw:frame draw:name="Link" svg:x="13cm" svg:y="5cm" svg:width="10cm" svg:height="2cm"><draw:text-box><text:p><text:a xlink:href="https://example.invalid/odp" xlink:type="simple">ODP reference</text:a></text:p></draw:text-box></draw:frame>'
        body+=f'<presentation:notes><draw:page-thumbnail draw:page-number="{index+1}" svg:x="1cm" svg:y="1cm" svg:width="10cm" svg:height="6cm"/><draw:frame presentation:class="notes" draw:name="Notes{index+1}" svg:x="1cm" svg:y="8cm" svg:width="17cm" svg:height="10cm"><draw:text-box><text:p>ODP-SPEAKER-NOTE-{index+1}</text:p></draw:text-box></draw:frame></presentation:notes></draw:page>'
    body+='</office:presentation>'
    path=GENERATED/'presentation-original-odf.odp'; write_odf(path,'odp',body,'',image)
    register(path,'presentation',['original-odf','three-slides','table','image','editable-shape','notes','hyperlink'],{'textIncludes':['COMPAT-ORIGINAL-ODP','ODP-DATA-SLIDE','ODP-FINAL-SLIDE'],'slideCount':3,'minimumFeatures':{'tables':1,'images':1,'hyperlinks':1}},'Original ODF 1.2 XML/ZIP generator; not LibreOffice-authored')

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true', help='Verify current fixture hashes without generating files')
    parser.add_argument('--preserve-existing', action='store_true', help='Generate only missing fixtures; preserve existing exact fixture bytes')
    args = parser.parse_args()
    if args.check:
        manifest = json.loads(MANIFEST.read_text())
        for item in manifest['fixtures']:
            payload = (ROOT / item['path']).read_bytes()
            assert len(payload) == item['bytes'] and hashlib.sha256(payload).hexdigest() == item['sha256'], item['id']
        print(f"Verified {len(manifest['fixtures'])} exact fixture hashes")
        return
    GENERATED.mkdir(parents=True, exist_ok=True)
    image = brand_image()
    previous = json.loads(MANIFEST.read_text()) if MANIFEST.exists() else {}
    steps = [(word_rich,'writer-rich'),(workbook_rich,'calc-rich'),(presentation_rich,'presentation-rich'),(word_long,'writer-long-report'),(word_review,'writer-review-footnotes'),(workbook_large,'calc-2000-rows'),(workbook_dates,'calc-dates-validation'),(presentation_long,'presentation-20-slides'),(presentation_effects,'presentation-effects-media'),(word_rtf,'writer-unicode-tables'),(word_odf,'writer-original-odf'),(workbook_odf,'calc-original-odf'),(presentation_odf,'presentation-original-odf')]
    existing = {item['id']:item for item in previous.get('fixtures',[])}
    for generate, ident in steps:
        if args.preserve_existing and ident in existing and (ROOT/existing[ident]['path']).exists():
            entry=existing[ident]
            if 'generatedAt' in entry['provenance']: entry['provenance']['reproducibleTimestamp']=entry['provenance'].pop('generatedAt')
            ENTRIES.append(entry)
        else:
            generate(image)
    upstream = [item for item in previous.get('fixtures', []) if item.get('provenance', {}).get('type') != 'generated']
    manifest = {
        'schemaVersion': 1,
        'description': 'Original generated and separately attributed upstream office compatibility fixtures; no claim of Microsoft/WPS interoperability certification.',
        'fixtureRoot': 'tests/compatibility',
        'fixtures': ENTRIES + upstream,
    }
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
    print(f'Generated {len(ENTRIES)} fixtures; retained {len(upstream)} upstream fixtures')


if __name__ == '__main__':
    main()
