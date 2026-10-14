from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from pathlib import Path
import json,hashlib
out=Path(__file__).parent
wb=Workbook(); ws=wb.active;ws.title='Mixed format'
for row in range(1,10):
 for col in range(1,4):
  c=ws.cell(row,col);c.value=1234.5 if col==1 else f'=A{row}*2' if col==2 else '00123'
  c.font=Font(name=['Calibri','Arial','Courier New'][col-1],size=[11,20,14][col-1],bold=col==1,italic=col==2,color=['FF990000','FF000099','FF008800'][col-1],underline='single' if col==3 else None,strike=col==3)
  c.fill=PatternFill('solid',fgColor=['FFFFCC00','FF00FFFF','FFFF00FF'][col-1]);c.alignment=Alignment(horizontal=['left','center','right'][col-1],vertical=['top','center','bottom'][col-1],wrap_text=col==1)
  c.border=Border(left=Side(style=['thin','double','dashed'][col-1],color='FFAA00FF'))
  c.number_format=['#,##0.00','0.0%','@'][col-1]
 ws.cell(row,4).value=CellRichText(TextBlock(InlineFont(rFont='Calibri',sz=12,b=True,color='FFFF0000'),'RED'),TextBlock(InlineFont(rFont='Arial',sz=20,i=True,color='FF0000FF'),'BLUE'))
 ws.cell(row,4).fill=PatternFill('solid',fgColor='FFCCCCCC');ws.cell(row,4).number_format='@'
ws.column_dimensions['A'].width=18;ws.column_dimensions['B'].width=22;ws.column_dimensions['C'].width=18;ws.column_dimensions['D'].width=24
file=out/'mixed-format-source.xlsx';wb.save(file)
(out/'fixture.json').write_text(json.dumps({'file':file.name,'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'generator':'openpyxl 3.1.5; independent cell and rich-run fonts','rows':9,'columns':4},indent=2))
