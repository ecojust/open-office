"""Original OOXML edge cases; python-pptx supplies only the package scaffold.

The explicitly written XML and independent expectations test the application
parser, not a round trip through that application's writer. MIT licensed.
"""
import hashlib
import json
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo
from io import BytesIO
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
FONT = '<a:latin typeface="Arial"/><a:ea typeface="Arial"/>'


def r(text, props=''):
    from xml.sax.saxutils import escape
    return f'<a:r><a:rPr {props}>{FONT}</a:rPr><a:t>{escape(text)}</a:t></a:r>'


def p(runs, props='', end='sz="2400"'):
    return f'<a:p><a:pPr {"" if props.startswith("<") else props}>{props if props.startswith("<") else ""}</a:pPr>{runs}<a:endParaRPr {end}>{FONT}</a:endParaRPr></a:p>'


def textbox(slide, name, x, y, width, height, paragraphs, list_style='', body_extra=''):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    shape.name = name
    body = etree.fromstring(f'<p:txBody xmlns:a="{A}" xmlns:p="{P}"><a:bodyPr wrap="none" anchor="t" lIns="0" rIns="0" tIns="0" bIns="0" {body_extra}/><a:lstStyle>{list_style}</a:lstStyle>{paragraphs}</p:txBody>')
    shape._element.replace(shape._element.find(f'{{{P}}}txBody'), body)
    return shape


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
props = prs.core_properties
props.title = 'Spacing and baseline - independent OOXML fixture'
props.author = 'Open Office compatibility tests'
props.created = props.modified = datetime(2025, 2, 12, 12, 0, 0)

s = prs.slides.add_slide(prs.slide_layouts[6])
textbox(s, 'TITLE_SPACING', .35, .25, 12, .6, p(r('Explicit paragraph spacing and line advances', 'sz="2600" b="1"')))
for x, name, label, spacing in [(.4, 'LINE_90', '90% line spacing', '<a:spcPct val="90000"/>'), (4.6, 'LINE_150', '150% line spacing', '<a:spcPct val="150000"/>'), (8.8, 'LINE_FIXED', 'Fixed 12pt (intentional overlap)', '<a:spcPts val="1200"/>')]:
    textbox(s, name+'_LABEL', x, 1.0, 4.0, .6, p(r(label, 'sz="1500" b="1"')))
    pp=f'<a:lnSpc>{spacing}</a:lnSpc><a:spcBef><a:spcPts val="600"/></a:spcBef><a:spcAft><a:spcPts val="800"/></a:spcAft><a:defRPr sz="2400">{FONT}</a:defRPr>'
    runs=r('First 24pt line','sz="2400"')+'<a:br/>'+r('Second 24pt line','sz="2400"')+'<a:br/>'+r('Third 24pt line','sz="2400"')
    textbox(s,name,x,1.8,4.0,4.8,p(runs,pp)+p(r('Next paragraph','sz="2400"'),pp))

s=prs.slides.add_slide(prs.slide_layouts[6])
textbox(s,'TITLE_BASELINE',.35,.25,12,.6,p(r('Explicit positive, negative and zero baseline', 'sz="2600" b="1"')))
colors=['202020','B51D46','1565A8','248046']
runs=''
for text, baseline, color in zip(['Normal ', 'Raised ', 'Lowered ', 'Zero'], [0,30000,-25000,0],colors):
    runs += r(text,f'sz="2600" baseline="{baseline}"').replace(FONT,f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'+FONT)
textbox(s,'BASELINE_EXPLICIT',.7,1.65,12,1.4,p(runs))
textbox(s,'BASELINE_FRACTIONAL',.7,3.6,12,1.4,p(r('Normal ', 'sz="2600"')+r('18 percent ', 'sz="2600" baseline="18000"')+r('minus 10 percent ', 'sz="2600" baseline="-10000"')+r('12.34pt','sz="1234" baseline="0"')))
textbox(s,'FOOTNOTE',.7,6,12,.8,p(r('Source numeric values must survive edits and repeated saves.', 'sz="1700"')))

s=prs.slides.add_slide(prs.slide_layouts[6])
textbox(s,'TITLE_INHERITANCE',.35,.25,12,.6,p(r('List defaults, explicit zero, empty paragraph', 'sz="2600" b="1"')))
defaults=f'<a:lvl1pPr><a:lnSpc><a:spcPct val="150000"/></a:lnSpc><a:spcBef><a:spcPts val="600"/></a:spcBef><a:spcAft><a:spcPts val="800"/></a:spcAft><a:defRPr sz="2400" baseline="30000">{FONT}</a:defRPr></a:lvl1pPr>'
first=p(r('Inherited raised ')+r('Explicit zero','baseline="0"'))
reset='<a:lnSpc><a:spcPct val="100000"/></a:lnSpc><a:spcBef><a:spcPts val="0"/></a:spcBef><a:spcAft><a:spcPts val="0"/></a:spcAft><a:defRPr baseline="0"/>'
second=p(r('Paragraph clears spacing and baseline')+'<a:br/>'+r('Manual break stays in this paragraph'),reset)
empty=p('',reset,'sz="3720" baseline="0"')
last=p(r('After empty paragraph','sz="2400" baseline="0"'))
textbox(s,'INHERITED',.7,1.4,12,5.7,first+second+empty+last,defaults)

s=prs.slides.add_slide(prs.slide_layouts[6])
textbox(s,'TITLE_BOUNDARIES',.35,.25,12,.6,p(r('Mixed-size fixed lines and first/last paragraph spacing', 'sz="2400" b="1"')))
textbox(s,'MIXED_FIXED_LABEL',.5,1,4,.5,p(r('12 / 42 / 18pt; fixed 12pt', 'sz="1500"')))
textbox(s,'MIXED_FIXED',.5,1.8,4,4,p(r('Small 12','sz="1200"')+'<a:br/>'+r('Large 42','sz="4200"')+'<a:br/>'+r('Medium 18','sz="1800"'),'<a:lnSpc><a:spcPts val="1200"/></a:lnSpc>'))
boundary='<a:lnSpc><a:spcPct val="100000"/></a:lnSpc><a:spcBef><a:spcPts val="1200"/></a:spcBef><a:spcAft><a:spcPts val="1600"/></a:spcAft>'
for x,flag in [(5,0),(9,1)]:
    textbox(s,f'OUTER_{flag}_LABEL',x,1,4,.5,p(r(f'spcFirstLastPara = {flag}', 'sz="1500"')))
    textbox(s,f'OUTER_{flag}',x,1.8,4,4,p(r('First paragraph','sz="2000"'),boundary)+p(r('Last paragraph','sz="2000"'),boundary),body_extra=f'spcFirstLastPara="{flag}"')

raw=BytesIO(); prs.save(raw)
path=ROOT/'spacing-baseline.pptx'
with ZipFile(raw) as src, ZipFile(path,'w',ZIP_DEFLATED) as dest:
    for name in sorted(src.namelist()):
        info=ZipInfo(name,(2025,2,12,12,0,0)); info.compress_type=ZIP_DEFLATED
        dest.writestr(info,src.read(name))
manifest={'schemaVersion':1,'items':[{'id':'spacing-baseline','path':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'expected':{'slides':len(prs.slides)},'provenance':{'kind':'generated','license':'MIT','generator':'generate.py','producerDeclared':'Original python-pptx scaffold plus explicit OOXML; not reference-editor generated'}}]}
(ROOT/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
print(json.dumps(manifest))
